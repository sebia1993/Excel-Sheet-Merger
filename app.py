"""PyInstaller entry point. Normal runs launch only the existing GUI.

An undocumented self-test switch is used by Windows CI to exercise the
*packaged executable* and all of its bundled spreadsheet readers.
"""
from __future__ import annotations

import sys
import tempfile
import traceback
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent / 'src'))


def _smoke_test(result_file: Path) -> None:
    from openpyxl import Workbook, load_workbook
    from excel_sheet_merger.classify import scan_folder
    from excel_sheet_merger.merge import merge_group

    with tempfile.TemporaryDirectory() as folder:
        directory = Path(folder)
        (directory / 'A사업장_20261028_11시.csv').write_text(
            '코드,IP\n00123,192.0.2.1\n', encoding='utf-8')
        src = directory / 'A사업장_20261028_14시.xlsx'
        wb = Workbook()
        wb.active['A1'] = 'Excel'
        wb.active['B1'] = '=1+2'
        wb.save(src)
        wb.close()
        report = scan_folder(directory)
        if len(report.groups) != 1:
            raise AssertionError('folder scan did not find exactly one merge group')
        result = merge_group(report.groups[0], directory/'병합결과')
        joined = load_workbook(result.path)
        try:
            assert joined.sheetnames == ['11시', '14시']
            assert joined['11시']['A2'].value == '00123'
            assert joined['14시']['B1'].value == '=1+2'
        finally:
            joined.close()
    result_file.write_text('PASS: CSV+XLSX merge and output verification\n', encoding='utf-8')


if __name__ == '__main__':
    if len(sys.argv) == 3 and sys.argv[1] == '--self-test':
        try:
            _smoke_test(Path(sys.argv[2]))
        except Exception:
            Path(sys.argv[2]).write_text('FAIL:\n' + traceback.format_exc(), encoding='utf-8')
            raise SystemExit(1)
    else:
        from excel_sheet_merger.gui import run_app
        run_app()
