"""GUI 작업 스레드와 Windows EXE 패키징 계약."""
import queue
import sys
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))


class DistributionTests(unittest.TestCase):
    def test_gui_import_does_not_need_excel_or_pywin32(self):
        from excel_sheet_merger.gui import run_app, summarize_scan
        self.assertTrue(callable(run_app))
        self.assertTrue(callable(summarize_scan))

    def test_merge_worker_uses_folder_path_without_tkinter_variable_access(self):
        from excel_sheet_merger.gui import MergeApp
        from excel_sheet_merger.classify import scan_folder
        with tempfile.TemporaryDirectory() as d:
            source = Path(d)
            for time in ("11시", "14시"):
                (source / f"A사업장_20261028_{time}.csv").write_text("ID,주소\n001,192.168.1.1\n", encoding="utf-8")
            group = scan_folder(source).groups[0]
            app = MergeApp.__new__(MergeApp)
            app.events = queue.Queue()

            class ForbiddenFolder:
                def get(self):
                    raise AssertionError("worker must not read Tk variables")

            app.folder = ForbiddenFolder()
            outcome = app._process_group(group, source / "병합결과", 1, 1, 0, 0, 0, [])
            self.assertEqual(outcome, (1, 0, 0))
            xlsx = source / "병합결과" / "A사업장_20261028.xlsx"
            wb = load_workbook(xlsx, read_only=True)
            try:
                self.assertEqual(wb.sheetnames, ["11시", "14시"])
            finally:
                wb.close()

    def test_worker_continues_after_corrupt_group(self):
        from excel_sheet_merger.gui import MergeApp
        from excel_sheet_merger.classify import scan_folder
        with tempfile.TemporaryDirectory() as d:
            p = Path(d)
            for site, t, payload in (
                ("A사업장", "11시", b"col\nnormal\n"),
                ("A사업장", "14시", b"\xff\x80\xff"),
                ("B사업장", "11시", b"col\nfirst\n"),
                ("B사업장", "14시", b"col\nsecond\n"),
            ):
                (p / f"{site}_20261028_{t}.csv").write_bytes(payload)
            scan = scan_folder(p)
            app = MergeApp.__new__(MergeApp)
            app.events = queue.Queue()
            app.report = scan
            app._worker_merge(list(scan.groups), p / "병합결과")
            events = list(app.events.queue)
            finished = [e for e in events if e[0] == "done"]
            self.assertEqual(finished, [("done", 1, 1, 0)])
            self.assertTrue((p / "병합결과" / "B사업장_20261028.xlsx").exists())
            self.assertFalse((p / "병합결과" / "A사업장_20261028.xlsx").exists())

    def test_windows_builder_packages_openpyxl_without_excel_com(self):
        workflow = (ROOT / ".github/workflows/windows-exe.yml").read_text(encoding="utf-8")
        bat = (ROOT / "build_windows.bat").read_text(encoding="utf-8")
        requirements = (ROOT / "requirements-build.txt").read_text(encoding="utf-8")
        self.assertIn("windows-latest", workflow)
        self.assertIn("pyinstaller", workflow.lower())
        self.assertIn("unittest discover", workflow)
        self.assertIn("upload-artifact", workflow)
        self.assertIn("openpyxl", requirements)
        self.assertNotIn("pywin32", requirements.lower())
        self.assertNotIn("win32com", workflow.lower() + bat.lower())
        self.assertIn("contents: write", workflow)
        self.assertIn("gh release create", workflow)
        self.assertIn("SHA256SUMS", workflow)
        # Standalone Windows EXE must include the themed ttkbootstrap assets.
        self.assertIn("ttkbootstrap==2.2.6", requirements)
        self.assertIn("--collect-all ttkbootstrap", workflow)
        self.assertIn("--collect-all ttkbootstrap", bat)

    def test_user_documentation_describes_csv_to_xlsx_conversion(self):
        readme = (ROOT / "README.md").read_text(encoding="utf-8")
        self.assertIn("A사업장_20261028_11시.csv", readme)
        self.assertIn("A사업장_20261028.xlsx", readme)
        self.assertIn("11시", readme)
        self.assertIn("Excel 설치가 필요하지", readme)
        self.assertIn("PyInstaller", readme)
        self.assertIn("releases/latest/download/ExcelSheetMerger.exe", readme)


if __name__ == "__main__":
    unittest.main()
