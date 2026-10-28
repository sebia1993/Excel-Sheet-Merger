"""Integration guardrails for ZIP sheet stitching and unchanged Excel behavior."""
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from excel_sheet_merger.merge import _stitch_csv_sheets, text_cell


class ZipStitchIntegrityTests(unittest.TestCase):
    def test_replacement_preserves_untouched_members_and_csv_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            base = Workbook()
            excel = base.active
            excel.title = '11시'
            excel['A1'] = '원래 Excel'
            excel['B2'] = '=1+2'
            excel['B2'].font = Font(bold=True)
            excel.merge_cells('C2:D2')
            excel['C2'] = '병합 셀'
            base.create_sheet('14시')
            base_path = d / 'base.xlsx'
            base.save(base_path)
            base.close()

            csv = Workbook(write_only=True)
            sheet = csv.create_sheet('14시')
            sheet.append([text_cell(sheet, 'ID'), text_cell(sheet, '설명')])
            sheet.append([text_cell(sheet, '00123'), text_cell(sheet, '=SUM(A1:A2)')])
            csv_path = d / 'csv.xlsx'
            csv.save(csv_path)
            csv.close()

            output = d / 'result.xlsx'
            with zipfile.ZipFile(base_path) as source:
                untouched = {i.filename: (i.CRC, i.file_size, i.compress_type, i.date_time)
                             for i in source.infolist() if i.filename != 'xl/worksheets/sheet2.xml'}
            with zipfile.ZipFile(csv_path) as source:
                stream_bytes = source.read('xl/worksheets/sheet1.xml')

            _stitch_csv_sheets(base_path, csv_path, {2: 1}, output)
            with zipfile.ZipFile(output) as zf:
                self.assertIsNone(zf.testzip())
                self.assertEqual(zf.read('xl/worksheets/sheet2.xml'), stream_bytes)
                self.assertEqual(len(zf.namelist()), len(set(zf.namelist())))
                self.assertEqual(untouched, {
                    i.filename: (i.CRC, i.file_size, i.compress_type, i.date_time)
                    for i in zf.infolist() if i.filename != 'xl/worksheets/sheet2.xml'
                })
            result = load_workbook(output)
            try:
                self.assertEqual(result.sheetnames, ['11시', '14시'])
                self.assertEqual(result['11시']['B2'].value, '=1+2')
                self.assertEqual(result['11시']['B2'].data_type, 'f')
                self.assertTrue(result['11시']['B2'].font.bold)
                self.assertEqual(str(next(iter(result['11시'].merged_cells.ranges))), 'C2:D2')
                self.assertEqual(result['14시']['A2'].value, '00123')
                self.assertEqual(result['14시']['B2'].value, '=SUM(A1:A2)')
                self.assertEqual(result['14시']['B2'].data_type, 's')
            finally:
                result.close()


if __name__ == '__main__':
    unittest.main()
