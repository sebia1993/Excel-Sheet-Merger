"""Synthetic network/SMB failures must not corrupt outputs or mislead users."""
from __future__ import annotations

import errno
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from excel_sheet_merger.classify import scan_folder
from excel_sheet_merger.merge import SourceFileError, merge_group


class NetworkFailureTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.source = Path(self.tmp.name)
        for site in ('A사업장', 'B사업장'):
            for hour in ('11시', '14시'):
                (self.source / f'{site}_20261028_{hour}.csv').write_text(
                    'ID,VALUE\n00123,192.0.2.1\n', encoding='utf-8')
        self.target = self.source / '병합결과'

    def test_network_failure_errno_and_winerror_are_distinguished_from_other_io(self):
        from excel_sheet_merger.network_io import is_network_io_error
        for name in ('ENETDOWN', 'ENETUNREACH', 'ENOTCONN', 'ECONNRESET', 'ECONNABORTED', 'EHOSTUNREACH', 'ETIMEDOUT'):
            if hasattr(errno, name):
                with self.subTest(errno=name):
                    self.assertTrue(is_network_io_error(OSError(getattr(errno, name), 'offline')))
        class WindowsShareFailure(OSError):
            def __init__(self, code: int):
                super().__init__(errno.EINVAL, 'Windows network share unavailable')
                self._share_error = code

            @property
            def winerror(self):
                return self._share_error

        for code in (53, 64, 67, 121, 1231, 1232):
            with self.subTest(winerror=code):
                self.assertTrue(is_network_io_error(WindowsShareFailure(code)))
        for exc in (PermissionError(errno.EACCES, 'denied'), OSError(errno.ENOSPC, 'disk full'),
                    OSError(errno.EIO, 'local drive IO'), FileNotFoundError(errno.ENOENT, 'missing')):
            with self.subTest(non_network=str(exc)):
                self.assertFalse(is_network_io_error(exc))

    def test_network_loss_during_source_fingerprint_aborts_the_entire_scan(self):
        # If fingerprint fails, classification must not accept an unverified
        # subset of the network directory as a mergeable report.
        with patch('excel_sheet_merger.classify.file_fingerprint',
                   side_effect=OSError(errno.ENETDOWN, 'network down')):
            with self.assertRaises(OSError) as captured:
                scan_folder(self.source)
        self.assertEqual(captured.exception.errno, errno.ENETDOWN)
        self.assertFalse(self.target.exists())

    def test_permission_only_fingerprint_still_keeps_source_for_normal_error_guidance(self):
        # Corporate DLP can block read but allow stat; don't confuse this with
        # a lost network or silently alter existing DLP semantics.
        with patch('excel_sheet_merger.classify.file_fingerprint',
                   side_effect=PermissionError(errno.EACCES, 'restricted')):
            report = scan_folder(self.source)
        self.assertEqual(len(report.groups), 2)
        self.assertEqual(len(report.rejected), 0)
        self.assertTrue(all(f.content_fingerprint is None for g in report.groups for f in g.inputs))

    def test_network_loss_during_csv_read_fails_one_group_and_preserves_unrelated_work(self):
        from excel_sheet_merger import merge as merger
        groups = scan_folder(self.source).groups
        original_csv_rows = merger.csv_rows

        def flaky_rows(path: Path):
            if path.name.startswith('A사업장') and path.name.endswith('14시.csv'):
                raise OSError(errno.ECONNRESET, 'share disconnected mid-read')
            yield from original_csv_rows(path)

        with patch.object(merger, 'csv_rows', flaky_rows):
            with self.assertRaises(SourceFileError) as captured:
                merge_group(groups[0], self.target)
            good = merge_group(groups[1], self.target)
        self.assertIn('14시.csv', captured.exception.filename)
        self.assertFalse((self.target / groups[0].output_name).exists())
        self.assertEqual(good.path.name, groups[1].output_name)
        self.assertEqual(list(self.target.glob('_병합임시_*')), [])
        workbook = load_workbook(good.path, read_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ['11시', '14시'])
        finally:
            workbook.close()

    def test_network_loss_during_final_archive_write_keeps_existing_output_and_cleans_scratch(self):
        from excel_sheet_merger import merge as merger
        from openpyxl import Workbook
        # Mixed CSV+XLSX group uses the final ZIP stitching step.
        excel = self.source / 'A사업장_20261028_14시.xlsx'
        (self.source / 'A사업장_20261028_14시.csv').unlink()
        w = Workbook(); w.active['A1'] = 'Excel'; w.save(excel); w.close()
        group = next(g for g in scan_folder(self.source).groups if g.site == 'A사업장')
        self.target.mkdir()
        old = self.target / group.output_name
        old.write_bytes(b'KEEP-ORIGINAL')
        with patch.object(merger, '_stitch_csv_sheets',
                          side_effect=OSError(errno.ENETUNREACH, 'share unreachable')):
            with self.assertRaises(OSError):
                merge_group(group, self.target)
        self.assertEqual(old.read_bytes(), b'KEEP-ORIGINAL')
        self.assertFalse((self.target / (old.stem + ' (2).xlsx')).exists())
        self.assertEqual(list(self.target.glob('_병합임시_*')), [])

    def test_publish_connection_lost_after_server_saved_result_is_reported_as_uncertain(self):
        from excel_sheet_merger import merge as merger
        from excel_sheet_merger.failure_guidance import diagnose_failure
        group = scan_folder(self.source).groups[0]
        actual_publish = os.rename if os.name == 'nt' else os.link
        method = 'rename' if os.name == 'nt' else 'link'
        calls = []

        def server_committed_but_reply_lost(staged, final):
            calls.append(Path(final))
            actual_publish(staged, final)
            raise OSError(errno.ETIMEDOUT, 'server did not respond after commit')

        with patch('excel_sheet_merger.merge.os.' + method, side_effect=server_committed_but_reply_lost):
            with self.assertRaises(OSError) as caught:
                merge_group(group, self.target)
        # The client cannot know whether a network timeout followed a commit.
        self.assertEqual(len(calls), 1, 'never automatically retry unknown publication outcome')
        self.assertTrue((self.target / group.output_name).exists())
        self.assertEqual(list(self.target.glob('_병합임시_*')), [])
        guidance = diagnose_failure(caught.exception, group)
        self.assertEqual(guidance.code, 'NETWORK_IO')
        self.assertIn('결과', guidance.action)
        self.assertIn('재시도', guidance.action)

    def test_network_diagnostic_does_not_mask_local_disk_or_permission_errors(self):
        from excel_sheet_merger.failure_guidance import diagnose_failure
        network = diagnose_failure(OSError(errno.ENETDOWN, 'network down'))
        self.assertEqual(network.code, 'NETWORK_IO')
        self.assertIn('공유폴더', network.reason)
        self.assertNotIn('암호화', network.reason)
        self.assertIn('연결', network.action)
        self.assertEqual(diagnose_failure(OSError(errno.ENOSPC, 'full')).code, 'DISK_FULL')
        self.assertEqual(diagnose_failure(PermissionError(errno.EACCES, 'denied')).code, 'PERMISSION')
        self.assertEqual(diagnose_failure(OSError(errno.EIO, 'local disk hardware')).code, 'UNKNOWN')

    def test_nested_network_failure_takes_precedence_over_wrapped_excel_open_message(self):
        from excel_sheet_merger.failure_guidance import diagnose_failure
        try:
            try:
                raise OSError(errno.ENETDOWN, 'share failed')
            except OSError as original:
                raise SourceFileError('A사업장_20261028_14시.xlsx', ValueError('Excel 파일을 열 수 없습니다')) from original
        except SourceFileError as wrapped:
            guidance = diagnose_failure(wrapped)
        self.assertEqual(guidance.filename, 'A사업장_20261028_14시.xlsx')
        self.assertEqual(guidance.code, 'NETWORK_IO')


@unittest.skipUnless(os.name == 'nt' or bool(os.environ.get('DISPLAY')),
                     'GUI requires Windows desktop or DISPLAY')
class NetworkScanGuiTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.source = Path(temp.name)
        for hour in ('11시', '14시'):
            (self.source / f'A사업장_20261028_{hour}.csv').write_text('ID\n00123\n', encoding='utf-8')

    def test_disconnected_share_invalidates_preview_and_shows_korean_recovery(self):
        from excel_sheet_merger.gui import MergeApp
        app = MergeApp()
        self.addCleanup(app.destroy)
        app.folder.set(str(self.source))
        app._scan()
        self.assertEqual(len(app.report.groups), 1)
        with patch('excel_sheet_merger.classify.file_fingerprint',
                   side_effect=OSError(errno.ENETDOWN, 'network down')), \
             patch('excel_sheet_merger.gui.messagebox.showerror') as popup:
            app._scan()
        self.assertIsNone(app.report)
        self.assertEqual(app.group_map, {})
        self.assertEqual(str(app.all_button.cget('state')), 'disabled')
        popup.assert_called_once()
        notice = popup.call_args.args[1]
        self.assertIn('네트워크', notice)
        self.assertIn('공유폴더', notice)
        self.assertIn('결과', notice)



if __name__ == '__main__':
    unittest.main()
