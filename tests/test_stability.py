"""Post-audit safety regression tests: actual behavior, not source-text assertions."""
from __future__ import annotations

import os
import sys
import queue
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook, Workbook

from excel_sheet_merger.classify import scan_folder
from excel_sheet_merger.merge import merge_group, SourceFileError
from excel_sheet_merger.gui import MergeApp


class _CaseBase(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.home = Path(self.temp.name)

    def input_folder(self, name='A'):
        folder = self.home / name
        folder.mkdir(exist_ok=True)
        for hour in ['11시', '14시']:
            (folder / f'사업장_20261028_{hour}.csv').write_text('IP,코드\n192.0.2.1,00123\n', encoding='utf-8')
        return folder

class StabilityTests(_CaseBase):
    def test_scan_failure_invalidates_old_groups_and_prevents_cross_folder_publish(self):
        a = self.input_folder('a')
        b = self.home / 'b'; b.mkdir()
        app = MergeApp()
        self.addCleanup(app.destroy)
        app.folder.set(str(a)); app._scan()
        self.assertEqual(len(app.report.groups), 1)
        app.folder.set(str(b))
        with patch('excel_sheet_merger.gui.scan_folder', side_effect=PermissionError('접근 거부')), \
             patch('excel_sheet_merger.gui.messagebox.showerror'):
            app._scan()
        self.assertIsNone(app.report)
        self.assertEqual(app.group_map, {})
        self.assertFalse(app.checked_group_ids)
        self.assertEqual(str(app.all_button['state']), 'disabled')
        self.assertEqual(app.tree.get_children(), ())
        app._start_merge(False)
        self.assertFalse((b / '병합결과').exists())
        app.folder.set(str(a)); app._scan()
        self.assertEqual(len(app.report.groups), 1)

    def test_scan_then_changed_input_fails_before_publishing(self):
        folder = self.input_folder()
        report = scan_folder(folder)
        source = folder / '사업장_20261028_11시.csv'
        source.write_text('ID\nMODIFIED DATA MANY BYTES\n', encoding='utf-8')
        with self.assertRaises(SourceFileError):
            merge_group(report.groups[0], folder / '병합결과')
        self.assertFalse(list((folder / '병합결과').glob('*.xlsx')))

    def test_scan_then_unchanged_source_still_merges_normally(self):
        folder = self.input_folder()
        result = merge_group(scan_folder(folder).groups[0], folder / '병합결과')
        wb = load_workbook(result.path, read_only=True)
        try:
            self.assertEqual(wb.sheetnames, ['11시', '14시'])
            self.assertEqual(wb['11시']['B2'].value, '00123')
        finally:
            wb.close()

if __name__ == '__main__':
    unittest.main()

class MemorySafetyTests(_CaseBase):
    def test_sparse_styled_sheet_does_not_traverse_blank_grid(self):
        folder = self.input_folder()
        excel = folder / '사업장_20261028_14시.xlsx'
        (folder / '사업장_20261028_14시.csv').unlink()
        wb = Workbook()
        ws = wb.active
        ws['A1'] = 'FIRST'
        ws['ZZ700'].font = __import__('openpyxl').styles.Font(bold=True)
        ws.row_dimensions[700].height = 33
        wb.save(excel); wb.close()
        from openpyxl.worksheet.worksheet import Worksheet
        with patch.object(Worksheet, 'iter_rows', side_effect=AssertionError('full sparse grid scan')):
            output = merge_group(scan_folder(folder).groups[0], folder / '병합결과')
        result = load_workbook(output.path)
        try:
            self.assertEqual(result['14시']['A1'].value, 'FIRST')
            self.assertTrue(result['14시']['ZZ700'].font.bold)
            self.assertEqual(result['14시'].row_dimensions[700].height, 33)
        finally:
            result.close()

    def test_mixed_csv_sheets_are_streamed_without_normal_worksheet_append(self):
        from openpyxl.worksheet.worksheet import Worksheet
        folder = self.home / 'mixed'; folder.mkdir()
        for hour in ['11시', '17시']:
            (folder / f'사업장_20261028_{hour}.csv').write_text('ID,메모\n00123,=SUM(A1:A3)\n', encoding='utf-8')
        book = Workbook(); sheet = book.active
        sheet['A1'] = 'Excel'; sheet['A2'] = '=1+1'
        sheet.merge_cells('B2:C2'); sheet['B2'] = '합쳐짐'
        sheet.column_dimensions['C'].width = 21
        sheet['A1'].font = __import__('openpyxl').styles.Font(bold=True)
        book.save(folder / '사업장_20261028_14시.xlsx'); book.close()
        base_append = Worksheet.append
        def no_normal_csv_append(ws, cells):
            if ws.title in ('11시','17시'):
                raise AssertionError('CSV data should stream, not use in-memory cells')
            return base_append(ws,cells)
        with patch.object(Worksheet, 'append', no_normal_csv_append):
            out = merge_group(scan_folder(folder).groups[0], folder/'병합결과')
        result = load_workbook(out.path)
        try:
            self.assertEqual(result.sheetnames,['11시','14시','17시'])
            for title in ('11시','17시'):
                self.assertEqual(result[title]['A2'].value,'00123')
                self.assertEqual(result[title]['B2'].data_type,'s')
            self.assertEqual(result['14시']['A2'].value,'=1+1')
            self.assertTrue(result['14시']['A1'].font.bold)
            self.assertEqual(str(next(iter(result['14시'].merged_cells.ranges))),'B2:C2')
            self.assertEqual(result['14시'].column_dimensions['C'].width,21)
        finally:
            result.close()

class ScratchRecoveryTests(_CaseBase):
    def test_old_abandoned_owned_scratch_removed_but_unknown_temp_is_preserved(self):
        from excel_sheet_merger.temp_workspace import OwnedScratch, cleanup_stale_workspaces
        folder = self.input_folder()
        dest = folder / '병합결과'
        dest.mkdir()
        with OwnedScratch(dest) as scratch:
            orphan = scratch.path
            (orphan / 'result.xlsx').write_bytes(b'INCOMPLETE')
            old = __import__('time').time() - 90000
            os.utime(orphan / 'owner.lock', (old, old))
        # Simulate force-kill residue: reclaim the expected files after the lock is gone.
        orphan.mkdir()
        (orphan / 'owner.lock').write_bytes(b'ExcelSheetMerger:scratch:v1\n')
        (orphan / 'result.xlsx').write_bytes(b'INCOMPLETE')
        os.utime(orphan / 'owner.lock', (old, old))
        unrelated = dest / '_병합중_11111111111111111111111111111111.xlsx'
        unrelated.write_bytes(b'UNKNOWN-OWNER')
        cleanup_stale_workspaces(dest)
        self.assertFalse(orphan.exists())
        self.assertEqual(unrelated.read_bytes(), b'UNKNOWN-OWNER')

    def test_active_locked_scratch_is_never_deleted_by_other_worker(self):
        from excel_sheet_merger.temp_workspace import OwnedScratch, cleanup_stale_workspaces
        dest = self.home / 'out'; dest.mkdir()
        with OwnedScratch(dest) as scratch:
            from excel_sheet_merger.temp_workspace import AGE_SECONDS
            with patch('excel_sheet_merger.temp_workspace.time.time',
                       return_value=__import__('time').time() + AGE_SECONDS + 60):
                cleanup_stale_workspaces(dest)
            self.assertTrue(scratch.path.exists())
        self.assertFalse(scratch.path.exists())

    def test_unknown_scratch_contents_are_preserved(self):
        from excel_sheet_merger.temp_workspace import cleanup_stale_workspaces
        dest = self.home / 'out'; dest.mkdir()
        folder = dest / ('_병합임시_' + '0'*32)
        folder.mkdir()
        marker = folder / 'owner.lock'
        marker.write_bytes(b'ExcelSheetMerger:scratch:v1\n')
        (folder / 'personnel_notes.txt').write_text('IMPORTANT')
        old = __import__('time').time()-90000
        os.utime(marker,(old,old))
        cleanup_stale_workspaces(dest)
        self.assertTrue(folder.exists())

    def test_normal_merge_does_not_leave_staging_directories(self):
        folder = self.input_folder()
        merge_group(scan_folder(folder).groups[0], folder/'병합결과')
        self.assertFalse(list((folder/'병합결과').glob('_병합임시_*')))

class GuiResilienceTests(_CaseBase):
    def test_log_write_failure_is_warning_without_losing_completed_result(self):
        folder = self.input_folder()
        app = MergeApp()
        self.addCleanup(app.destroy)
        app.folder.set(str(folder)); app._scan()
        result_dir = folder / '병합결과'
        with patch.object(app, '_write_job_record', side_effect=OSError(28, 'No space left')), \
             patch('excel_sheet_merger.gui.messagebox.showwarning') as warning, \
             patch('excel_sheet_merger.gui.messagebox.showinfo') as info:
            app._worker_merge(list(app.report.groups), result_dir)
            app._consume_events()
        self.assertTrue((result_dir/'사업장_20261028.xlsx').is_file())
        self.assertIn('완료 1개', app.state_text.get())
        self.assertTrue(warning.called)
        self.assertFalse(info.called)
        self.assertIn('작업 기록', warning.call_args.args[1])
        self.assertFalse(list(result_dir.glob('병합작업기록_*.txt')))

    def test_batched_event_dispatch_preserves_event_order_and_done(self):
        folder = self.input_folder()
        app = MergeApp()
        self.addCleanup(app.destroy)
        app.folder.set(str(folder)); app._scan()
        for i in range(240):
            app.events.put(('log', f'event-{i:03}'))
        app.events.put(('done',1,0,0))
        with patch('excel_sheet_merger.gui.messagebox.showinfo'):
            app._consume_events()
            self.assertGreater(app.events.qsize(),0, 'one callback must have a finite event budget')
            for _ in range(5):
                if app.events.empty(): break
                app._consume_events()
        self.assertTrue(app.events.empty())
        text=app.log.get('1.0','end')
        self.assertLess(text.index('event-001'), text.index('event-239'))
        self.assertEqual(app.state_text.get(), '완료 1개 / 이름 변경 0개 / 실패 0개')

    def test_changing_source_during_streaming_never_publishes_result(self):
        from excel_sheet_merger import merge as mod
        folder = self.input_folder()
        group = scan_folder(folder).groups[0]
        original = mod._append_csv
        changed = [False]
        def mutate_after_read(ws, path):
            original(ws,path)
            if not changed[0]:
                path.write_text('ID\nTHIS CHANGED WHILE READING\n',encoding='utf-8')
                changed[0] = True
        with patch.object(mod,'_append_csv',side_effect=mutate_after_read):
            with self.assertRaises(SourceFileError):
                merge_group(group,folder/'병합결과')
        self.assertFalse(list((folder/'병합결과').glob('*.xlsx')))

class PackagingSmokeTests(_CaseBase):
    def test_headless_packaged_smoke_command_exercises_csv_and_excel(self):
        import subprocess
        root=Path(__file__).resolve().parents[1]
        report=self.home/'smoke-result.txt'
        process=subprocess.run([sys.executable, str(root/'app.py'), '--self-test', str(report)],
                               capture_output=True,text=True,timeout=5)
        self.assertEqual(process.returncode,0,process.stdout+process.stderr)
        self.assertTrue(report.is_file())
        self.assertIn('PASS', report.read_text(encoding='utf-8'))
