"""Guard against source modifications that preserve filesystem metadata."""
from __future__ import annotations

import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))

from openpyxl import Workbook, load_workbook
from excel_sheet_merger.classify import scan_folder, file_fingerprint
from excel_sheet_merger.merge import SourceFileError, merge_group, _verify_source
from excel_sheet_merger.failure_guidance import diagnose_failure


class ContentFingerprintTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)

    def csv(self, hour, data):
        p = self.root / f'A사업장_20261028_{hour}.csv'
        p.write_bytes(data)
        return p

    @staticmethod
    def mutate_with_preserved_metadata(path: Path, replacement: bytes, offset: int = 0):
        before = path.stat()
        with path.open('r+b') as output:
            output.seek(offset)
            output.write(replacement)
            output.flush()
            os.fsync(output.fileno())
        os.utime(path, ns=(before.st_atime_ns, before.st_mtime_ns))
        after = path.stat()
        assert (after.st_size, after.st_mtime_ns, after.st_dev, after.st_ino) == (
            before.st_size, before.st_mtime_ns, before.st_dev, before.st_ino
        )

    def test_same_size_same_mtime_mutated_csv_is_not_published(self):
        source = self.csv('11시', b'ID,VALUE\n0001,HELLO\n')
        self.csv('14시', b'ID,VALUE\n0002,WORLD\n')
        report = scan_folder(self.root)
        self.mutate_with_preserved_metadata(source, b'9999', offset=9)
        with self.assertRaisesRegex(SourceFileError, '변경'):
            merge_group(report.groups[0], self.root / '병합결과')
        self.assertFalse(list((self.root / '병합결과').glob('*.xlsx')))

    def test_large_source_mutated_near_middle_with_same_metadata_is_detected(self):
        big = self.csv('11시', b'A' * (9 * 1024 * 1024))
        self.csv('14시', b'ID\n1\n')
        report = scan_folder(self.root)
        source = report.groups[0].inputs[0]
        self.mutate_with_preserved_metadata(big, b'REPLACEMENT', offset=big.stat().st_size // 2)
        with self.assertRaisesRegex(SourceFileError, '변경'):
            _verify_source(source)

    def test_large_source_changed_at_25_percent_with_identical_metadata_is_rejected(self):
        big = self.root / 'A사업장_20261028_11시.csv'
        with big.open('wb') as dest:
            dest.truncate(17 * 1024 * 1024)
        self.csv('14시', b'ID\n1\n')
        source = scan_folder(self.root).groups[0].inputs[0]
        self.mutate_with_preserved_metadata(big, b'ALTERED!', offset=big.stat().st_size // 4)
        with self.assertRaisesRegex(SourceFileError, '변경'):
            _verify_source(source)

    def test_large_source_changed_at_75_percent_with_identical_metadata_is_rejected(self):
        big = self.root / 'A사업장_20261028_11시.csv'
        with big.open('wb') as dest:
            dest.truncate(17 * 1024 * 1024)
        self.csv('14시', b'ID\n1\n')
        source = scan_folder(self.root).groups[0].inputs[0]
        self.mutate_with_preserved_metadata(big, b'ALTERED!', offset=big.stat().st_size * 3 // 4)
        with self.assertRaisesRegex(SourceFileError, '변경'):
            _verify_source(source)

    def test_confirmed_local_drive_detects_stealth_edit_between_old_samples(self):
        # Previous 17-window layout has a gap from 64 KiB to ~1.06 MiB.
        big = self.root / 'A사업장_20261028_11시.csv'
        with big.open('wb') as dest:
            dest.truncate(17 * 1024 * 1024)
        self.csv('14시', b'ID\n1\n')
        with patch('excel_sheet_merger.classify._is_fixed_local_drive',
                   return_value=True, create=True):
            scanned = scan_folder(self.root)
            self.mutate_with_preserved_metadata(big, b'INSIDE-GAP', offset=512 * 1024)
            with self.assertRaisesRegex(SourceFileError, '변경'):
                _verify_source(scanned.groups[0].inputs[0])

    def test_unconfirmed_local_drive_keeps_original_sample_budget(self):
        # A Windows mapped drive, UNC share, or unknown drive must not be
        # forced to stream the entire large file just for a fingerprint.
        source = self.root / 'large.bin'
        with source.open('wb') as dest:
            dest.truncate(17 * 1024 * 1024)
        actual_open = Path.open
        counted = []

        class MeasuredFile:
            def __init__(self, handle):
                self.handle = handle
            def __enter__(self):
                return self
            def __exit__(self, *args):
                self.handle.close()
            def read(self, size):
                contents = self.handle.read(size)
                counted.append(len(contents))
                return contents
            def seek(self, offset):
                return self.handle.seek(offset)

        def observed_open(path, mode='r', *args, **kwargs):
            handle = actual_open(path, mode, *args, **kwargs)
            if path == source and mode == 'rb':
                return MeasuredFile(handle)
            return handle

        with patch('excel_sheet_merger.classify._is_fixed_local_drive',
                   return_value=False, create=True), patch.object(Path, 'open', observed_open):
            fingerprint = file_fingerprint(source, source.stat().st_size)
        self.assertEqual(len(fingerprint), 20)
        self.assertEqual(counted, [64 * 1024] * 17)

    @unittest.skipUnless(os.name == 'nt', 'fixed/removable/remote Windows drive classification')
    def test_windows_local_disk_versus_unc_path(self):
        from excel_sheet_merger.classify import _is_fixed_local_drive
        self.assertTrue(_is_fixed_local_drive(self.root), 'runner temp directory is local NTFS')
        self.assertFalse(_is_fixed_local_drive(Path(r'\\server\share\folder\file.csv')))

    @unittest.skipUnless(os.name == 'nt', 'Windows drive-type API is only available on Windows')
    def test_windows_mapped_and_unknown_drives_are_not_assumed_fixed(self):
        from excel_sheet_merger.classify import _is_fixed_local_drive
        with patch('ctypes.windll.kernel32.GetDriveTypeW', return_value=4):
            self.assertFalse(_is_fixed_local_drive(self.root))  # DRIVE_REMOTE
        with patch('ctypes.windll.kernel32.GetDriveTypeW', return_value=1):
            self.assertFalse(_is_fixed_local_drive(self.root))  # DRIVE_NO_ROOT_DIR
        with patch('ctypes.windll.kernel32.GetDriveTypeW', side_effect=OSError('unavailable')):
            self.assertFalse(_is_fixed_local_drive(self.root))

    def test_large_file_fingerprint_uses_exactly_17_bounded_reads(self):
        source = self.root / 'large.bin'
        with source.open('wb') as dest:
            dest.truncate(64 * 1024 * 1024)
        byte_counts = []
        real_path = source

        class ObservedFile:
            def __init__(self):
                self.stream = real_path.open('rb')

            def __enter__(self):
                return self

            def __exit__(self, *_):
                self.stream.close()

            def seek(self, offset):
                return self.stream.seek(offset)

            def read(self, count):
                data = self.stream.read(count)
                byte_counts.append(len(data))
                return data

        class ObservedPath:
            def open(self, mode):
                assert mode == 'rb'
                return ObservedFile()

        fingerprint = file_fingerprint(ObservedPath(), source.stat().st_size)
        self.assertEqual(len(fingerprint), 20)
        self.assertEqual(byte_counts, [65536] * 17)
        self.assertEqual(sum(byte_counts), 17 * 65536)

    def test_full_fingerprint_rejects_unexpected_early_eof(self):
        source = self.root / 'short.bin'
        source.write_bytes(b'01234567')
        # An SMB metadata cache can report a size before the stream is cut.
        with self.assertRaisesRegex(ValueError, '원본 파일 읽기 불완전'):
            file_fingerprint(source, 100)

    def test_sampled_fingerprint_rejects_missing_window_bytes(self):
        source = self.root / 'short-large.bin'
        source.write_bytes(b'A' * (64 * 1024))
        with patch('excel_sheet_merger.classify._is_fixed_local_drive', return_value=False):
            with self.assertRaisesRegex(ValueError, '원본 파일 읽기 불완전'):
                file_fingerprint(source, 9 * 1024 * 1024)

    def test_partial_sample_reads_reassembled_without_digest_change(self):
        source = self.root / 'large.bin'
        with source.open('wb') as dest:
            dest.truncate(9 * 1024 * 1024)
        with source.open('r+b') as dest:
            dest.seek(8 * 1024 * 1024)
            dest.write(b'BEGIN-TEST-DATA')
        expected = file_fingerprint(source, source.stat().st_size)
        real_open = Path.open
        read_lengths = []

        class PartialReadWrapper:
            def __init__(self, handle):
                self.handle = handle
            def __enter__(self):
                return self
            def __exit__(self, *args):
                self.handle.close()
            def seek(self, offset):
                return self.handle.seek(offset)
            def read(self, count):
                data = self.handle.read(min(count, 32768))
                read_lengths.append(len(data))
                return data

        class NetworkStylePath:
            def open(self, mode):
                self.assert_mode(mode)
                return PartialReadWrapper(real_open(source, mode))
            @staticmethod
            def assert_mode(mode):
                assert mode == 'rb'

        with patch('excel_sheet_merger.classify._is_fixed_local_drive', return_value=False):
            actual = file_fingerprint(NetworkStylePath(), source.stat().st_size)
        self.assertEqual(actual, expected)
        self.assertEqual(sum(read_lengths), 17 * 65536)

    def test_full_fingerprint_rejects_unexpected_file_growth(self):
        source = self.root / 'growing.bin'
        source.write_bytes(b'ABCDEFGHIJ')
        with self.assertRaisesRegex(ValueError, '원본 파일 읽기 불완전'):
            file_fingerprint(source, 7)

    def test_verify_source_identifies_filename_when_fingerprint_read_is_incomplete(self):
        self.csv('11시', b'ID\n1\n')
        self.csv('14시', b'ID\n2\n')
        original = scan_folder(self.root).groups[0].inputs[0]
        with patch('excel_sheet_merger.merge.file_fingerprint',
                   side_effect=ValueError('원본 파일 읽기 불완전: 불완전한 읽기')):
            with self.assertRaises(SourceFileError) as caught:
                _verify_source(original)
        self.assertEqual(caught.exception.filename, original.file_name)
        self.assertIn('원본 파일 읽기 불완전', str(caught.exception))

    def test_incomplete_read_error_is_actionable_in_korean(self):
        filename = 'A사업장_20261028_11시.csv'
        advice = diagnose_failure(SourceFileError(
            filename, ValueError('원본 파일 읽기 불완전: 1바이트가 누락됨')))
        self.assertEqual(advice.filename, filename)
        self.assertEqual(advice.code, 'SOURCE_CHANGED')
        self.assertIn('다시', advice.action)
        self.assertIn('읽', advice.reason)

    def test_scan_does_not_accept_incomplete_content_fingerprint(self):
        short = self.csv('11시', b'ID\n1\n')
        self.csv('14시', b'ID\n2\n')
        from excel_sheet_merger.classify import file_snapshot
        real_snapshot = file_snapshot

        def stale_snapshot(path):
            original = real_snapshot(path)
            if path == short:
                return (original[0] + 65536, original[1], original[2], original[3])
            return original

        with patch('excel_sheet_merger.classify.file_snapshot', side_effect=stale_snapshot):
            report = scan_folder(self.root)
        self.assertEqual(report.groups, ())
        self.assertTrue(any(item.name == short.name for item in report.rejected))
        self.assertTrue(any('읽기 불완전' in item.reason for item in report.rejected))

    def test_unchanged_mixed_inputs_keep_sheet_names_values_and_formulas(self):
        self.csv('11시', 'IP,코드\n192.0.2.5,00123\n'.encode('utf-8'))
        wb = Workbook()
        wb.active['A1'] = '원본'
        wb.active['B2'] = '=1+2'
        wb.save(self.root / 'A사업장_20261028_14시.xlsx')
        wb.close()
        merged = merge_group(scan_folder(self.root).groups[0], self.root / '병합결과')
        result = load_workbook(merged.path)
        try:
            self.assertEqual(result.sheetnames, ['11시', '14시'])
            self.assertEqual(result['11시']['B2'].value, '00123')
            self.assertEqual(result['14시']['B2'].value, '=1+2')
        finally:
            result.close()

    def test_changed_source_error_explains_rescan_in_korean(self):
        error = SourceFileError('A사업장_20261028_11시.csv', ValueError(
            '파일 분석 이후 원본 파일 내용이 변경되었습니다. 다시 파일 분석을 실행하세요.'))
        advice = diagnose_failure(error)
        self.assertEqual(advice.code, 'SOURCE_CHANGED')
        self.assertEqual(advice.filename, 'A사업장_20261028_11시.csv')
        self.assertIn('변경', advice.reason)
        self.assertIn('파일 분석', advice.action)

    def test_scan_with_unreadable_fingerprint_keeps_group_for_existing_merge_errors(self):
        self.csv('11시', b'col\n1\n')
        self.csv('14시', b'col\n2\n')
        # Corporate file security may permit stat() but reject byte reads.
        # A new fingerprint check must not silently discard an otherwise
        # validly named source at the analysis stage.
        with patch('excel_sheet_merger.classify.file_fingerprint', side_effect=PermissionError('restricted')):
            report = scan_folder(self.root)
        self.assertEqual(len(report.groups), 1)
        self.assertTrue(all(item.content_fingerprint is None for item in report.groups[0].inputs))
        output = merge_group(report.groups[0], self.root / '병합결과')
        self.assertTrue(output.path.is_file())


if __name__ == '__main__':
    unittest.main()
