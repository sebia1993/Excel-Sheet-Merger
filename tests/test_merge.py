"""실제 CSV bytes를 XLSX로 변환하는 경계 테스트 (Excel COM 모의객체 없음)."""

import tempfile
import subprocess
import unittest
from pathlib import Path
from unittest.mock import patch
import sys

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from excel_sheet_merger.classify import scan_folder
from excel_sheet_merger.merge import MergeStatus, merge_group


class MergeTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.folder = Path(self.tmp.name)
        self.output = self.folder / "병합결과"

    def csv(self, time: str, contents: str, encoding: str = "utf-8", site: str = "A사업장") -> Path:
        path = self.folder / f"{site}_20261028_{time}.csv"
        path.write_bytes(contents.encode(encoding))
        return path

    def group(self):
        return scan_folder(self.folder).groups[0]

    def test_cp949_utf8_bom_and_multiline_values_to_three_time_sheets(self):
        source_11 = self.csv("11시", 'ID,메모,IP,MAC\r\n00123,"내용, 쉼표",10.0.0.1,00:11:22:33:44:55\r\n', "cp949")
        source_14 = self.csv("14시", 'ID,메모\r\n0001,"첫째 줄\r\n둘째 줄"\r\n', "utf-8-sig")
        source_17 = self.csv("17시", '식별자,원문\n900,=SUM(A1:A3)\n901,+3\n', "utf-8")
        original = [p.read_bytes() for p in (source_11, source_14, source_17)]
        result = merge_group(self.group(), self.output)
        self.assertEqual(result.status, MergeStatus.CREATED)
        self.assertEqual(result.path.name, "A사업장_20261028.xlsx")
        wb = load_workbook(result.path, read_only=False, data_only=False)
        try:
            self.assertEqual(wb.sheetnames, ["11시", "14시", "17시"])
            self.assertEqual([c.value for c in wb["11시"][2]], ["00123", "내용, 쉼표", "10.0.0.1", "00:11:22:33:44:55"])
            self.assertEqual(wb["14시"]["B2"].value, "첫째 줄\n둘째 줄")
            self.assertEqual(wb["17시"]["B2"].value, "=SUM(A1:A3)")
            self.assertEqual(wb["17시"]["B2"].data_type, "s")
            self.assertEqual(wb["17시"]["B3"].value, "+3")
            self.assertEqual(wb["11시"]["A2"].data_type, "s")
        finally:
            wb.close()
        self.assertEqual([p.read_bytes() for p in (source_11, source_14, source_17)], original)

    def test_duplicate_hours_produce_distinct_sheet_names(self):
        self.csv("11시", "값\n하나\n")
        self.csv("11-00", "값\n둘\n")
        self.csv("17시", "값\n셋\n")
        result = merge_group(self.group(), self.output)
        wb = load_workbook(result.path, read_only=True)
        try:
            self.assertEqual(wb.sheetnames, ["11시", "11시 (2)", "17시"])
            self.assertEqual([wb[s]["A2"].value for s in wb.sheetnames], ["둘", "하나", "셋"])
        finally:
            wb.close()

    def test_existing_workbook_gets_suffix_and_is_never_overwritten(self):
        self.csv("11시", "a\nb\n")
        self.csv("14시", "a\nc\n")
        self.output.mkdir()
        output = self.output / "A사업장_20261028.xlsx"
        output.write_bytes(b"EXISTING")
        result = merge_group(self.group(), self.output)
        self.assertEqual(result.status, MergeStatus.RENAMED)
        self.assertEqual(result.path.name, "A사업장_20261028 (2).xlsx")
        self.assertEqual(output.read_bytes(), b"EXISTING")
        wb = load_workbook(result.path, read_only=True)
        try:
            self.assertEqual(wb.sheetnames, ["11시", "14시"])
            self.assertEqual(wb["14시"]["A2"].value, "c")
        finally:
            wb.close()

    def test_repeated_merges_create_successive_numbered_files(self):
        self.csv("11시", "ID\n00123\n")
        self.csv("14시", "ID\n00456\n")
        group = self.group()
        first = merge_group(group, self.output)
        original = first.path.read_bytes()
        second = merge_group(group, self.output)
        third = merge_group(group, self.output)
        self.assertEqual([p.path.name for p in (first, second, third)], [
            "A사업장_20261028.xlsx",
            "A사업장_20261028 (2).xlsx",
            "A사업장_20261028 (3).xlsx",
        ])
        self.assertEqual([p.status for p in (first, second, third)], [
            MergeStatus.CREATED, MergeStatus.RENAMED, MergeStatus.RENAMED,
        ])
        self.assertEqual(first.path.read_bytes(), original)
        for created in (second, third):
            w = load_workbook(created.path, read_only=True)
            try:
                self.assertEqual(w.sheetnames, ["11시", "14시"])
                self.assertEqual(w["11시"]["A2"].value, "00123")
            finally:
                w.close()

    def test_atomic_publish_retries_after_another_writer_claims_the_name(self):
        self.csv("11시", "a\nb\n")
        self.csv("14시", "a\nc\n")
        from excel_sheet_merger.merge import os as merge_os
        publisher = merge_os.rename if merge_os.name == "nt" else merge_os.link
        function_name = "rename" if merge_os.name == "nt" else "link"
        output = self.output / "A사업장_20261028.xlsx"
        called = []

        def racing_writer(temp, destination):
            called.append(Path(destination).name)
            if Path(destination) == output:
                output.write_bytes(b"ANOTHER-WRITER")
                raise FileExistsError(17, "already exists", str(output))
            return publisher(temp, destination)

        with patch("excel_sheet_merger.merge.os." + function_name, side_effect=racing_writer):
            result = merge_group(self.group(), self.output)
        self.assertEqual(result.path.name, "A사업장_20261028 (2).xlsx")
        self.assertEqual(result.status, MergeStatus.RENAMED)
        self.assertEqual(called, ["A사업장_20261028.xlsx", "A사업장_20261028 (2).xlsx"])
        self.assertEqual(output.read_bytes(), b"ANOTHER-WRITER")
        self.assertTrue(result.path.is_file())

    def test_bad_source_does_not_write_suffix_when_existing_result_present(self):
        self.csv("11시", "a\nb\n")
        self.csv("14시", "")
        self.output.mkdir()
        existing = self.output / "A사업장_20261028.xlsx"
        existing.write_bytes(b"PREVIOUS")
        with self.assertRaisesRegex(ValueError, "비어"):
            merge_group(self.group(), self.output)
        self.assertEqual(existing.read_bytes(), b"PREVIOUS")
        self.assertEqual(sorted(p.name for p in self.output.iterdir()), ["A사업장_20261028.xlsx"])

    def test_unreadable_csv_fails_without_partial_workbook(self):
        self.csv("11시", "a\nb\n")
        damaged = self.folder / "A사업장_20261028_14시.csv"
        damaged.write_bytes(b"\xff\xff\x80\x80\n")
        with self.assertRaises((UnicodeError, ValueError, RuntimeError)):
            merge_group(self.group(), self.output)
        self.assertEqual(list(self.output.glob("*.xlsx")), [])
        self.assertEqual(list(self.output.glob("_병합중_*")), [])

    def test_empty_csv_fails_atomically(self):
        self.csv("11시", "정상\n")
        self.csv("14시", "")
        with self.assertRaisesRegex(ValueError, "비어"):
            merge_group(self.group(), self.output)
        self.assertEqual(list(self.output.glob("*.xlsx")), [])

    def test_quoted_csv_with_missing_quote_is_rejected(self):
        self.csv("11시", "a,b\n1,2\n")
        self.csv("14시", 'a,b\n1,"not closed\n')
        with self.assertRaises((ValueError, RuntimeError)):
            merge_group(self.group(), self.output)
        self.assertEqual(list(self.output.glob("*.xlsx")), [])

    def test_csv_excel_row_limit_is_enforced(self):
        self.csv("11시", "A\n1\n2\n")
        self.csv("14시", "A\n3\n")
        with patch("excel_sheet_merger.merge.MAX_ROWS", 2):
            with self.assertRaisesRegex(ValueError, "행"):
                merge_group(self.group(), self.output)
        self.assertEqual(list(self.output.glob("*.xlsx")), [])

    def test_csv_cell_length_limit_is_enforced_without_truncation(self):
        self.csv("11시", "a\n" + "가" * 32768 + "\n")
        self.csv("14시", "a\nb\n")
        with self.assertRaisesRegex(ValueError, "셀"):
            merge_group(self.group(), self.output)
        self.assertEqual(list(self.output.glob("*.xlsx")), [])

    def test_utf16_bom_csv_and_sep_directive(self):
        self.csv("11시", 'sep=,\r\nID,메모\r\n007,한글\r\n', "utf-16")
        self.csv("14시", 'ID,메모\r\n009,정상\r\n', "utf-8")
        result = merge_group(self.group(), self.output)
        wb = load_workbook(result.path)
        try:
            self.assertEqual(wb.sheetnames, ["11시", "14시"])
            self.assertEqual(wb["11시"]["A1"].value, "ID")
            self.assertEqual(wb["11시"]["A2"].value, "007")
            self.assertEqual(wb["11시"]["B2"].value, "한글")
        finally:
            wb.close()

    def test_failed_group_closes_streaming_writer_without_gc_errors(self):
        self.csv("11시", "header\nnormal\n")
        self.csv("14시", "")
        project_src = str(Path(__file__).resolve().parents[1] / "src")
        script = (
            "import sys,gc; sys.path.insert(0, sys.argv[1]); "
            "from pathlib import Path; "
            "from excel_sheet_merger.classify import scan_folder; "
            "from excel_sheet_merger.merge import merge_group; "
            "g=scan_folder(Path(sys.argv[2])).groups[0]; "
            "\ntry: merge_group(g,Path(sys.argv[2])/'병합결과')"
            "\nexcept ValueError: pass"
            "\ngc.collect()"
        )
        proc = subprocess.run(
            [sys.executable, "-c", script, project_src, str(self.folder)],
            capture_output=True, text=True, encoding="utf-8", timeout=20,
        )
        self.assertEqual(proc.returncode, 0, proc.stderr)
        self.assertNotIn("Exception ignored", proc.stderr)

    def test_csv_null_character_rejected_not_stripped(self):
        self.csv("11시", "a,b\n1,2\n")
        self.csv("14시", "a,b\n1,\x00\n")
        with self.assertRaisesRegex(ValueError, "문자"):
            merge_group(self.group(), self.output)
        self.assertFalse((self.output / "A사업장_20261028.xlsx").exists())


if __name__ == "__main__":
    unittest.main()
