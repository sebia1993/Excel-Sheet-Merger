"""Repeat-run regression: numbered outputs remain complete; scratch is reclaimed.

Uses only synthetic data and runs on Windows as well as Linux. This is a
small finite test, NOT a replacement for an eight-hour corporate-PC soak.
"""
from __future__ import annotations

import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from excel_sheet_merger.classify import scan_folder
from excel_sheet_merger.merge import MergeStatus, merge_group


class RepeatedMergeTests(unittest.TestCase):
    def test_100_successive_mixed_merges_keep_previous_versions_intact(self):
        with tempfile.TemporaryDirectory() as td:
            folder = Path(td)
            (folder / 'A사업장_20261028_11시.csv').write_text(
                'ID,주소,메모\n00123,192.0.2.1,"한글, 줄바꿈\n다음 줄"\n', encoding='utf-8')
            source = folder / 'A사업장_20261028_14시.xlsx'
            wb = Workbook()
            ws = wb.active
            ws['A1'] = '엑셀'
            ws['B1'] = '=1+2'
            ws['B1'].font = Font(bold=True)
            wb.save(source)
            wb.close()

            group = scan_folder(folder).groups[0]
            output_dir = folder / '병합결과'
            names = []
            original_hash = None
            for n in range(1, 101):
                result = merge_group(group, output_dir)
                names.append(result.path.name)
                self.assertEqual(result.status, MergeStatus.CREATED if n == 1 else MergeStatus.RENAMED)
                self.assertTrue(result.path.is_file())
                with zipfile.ZipFile(result.path) as zip_result:
                    self.assertIsNone(zip_result.testzip(), "every published XLSX must pass CRC checks")
                self.assertEqual(list(output_dir.glob('_병합임시_*')), [], 'no leaked working files')
                if n == 1:
                    original_hash = result.path.read_bytes()
                if n % 10 == 0 or n == 1:
                    opened = load_workbook(result.path, read_only=False)
                    try:
                        self.assertEqual(opened.sheetnames, ['11시', '14시'])
                        self.assertEqual(opened['11시']['A2'].value, '00123')
                        self.assertEqual(opened['11시']['C2'].value, '한글, 줄바꿈\n다음 줄')
                        self.assertEqual(opened['14시']['B1'].value, '=1+2')
                        self.assertTrue(opened['14시']['B1'].font.bold)
                    finally:
                        opened.close()
            self.assertEqual(names[0], 'A사업장_20261028.xlsx')
            self.assertEqual(names[1], 'A사업장_20261028 (2).xlsx')
            self.assertEqual(names[-1], 'A사업장_20261028 (100).xlsx')
            self.assertEqual(len(names), len(set(names)))
            self.assertEqual(len(list(output_dir.glob('*.xlsx'))), 100)
            self.assertEqual((output_dir / names[0]).read_bytes(), original_hash,
                             'the original result cannot be overwritten by a later run')


if __name__ == '__main__':
    unittest.main()
