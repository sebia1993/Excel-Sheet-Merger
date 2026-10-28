"""기존 CSV 변환과 함께 실제 Excel 입력을 병합하는 외부 관찰 테스트."""
import sys
import tempfile
import types
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
from excel_sheet_merger.classify import scan_folder
from excel_sheet_merger.merge import MergeStatus, merge_group


class ExcelInputTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.out = self.folder / '병합결과'

    def csv(self, hour, content):
        p = self.folder / f'A사업장_20261028_{hour}.csv'
        p.write_bytes(content.encode('utf-8'))
        return p

    def excel(self, hour, extension='.xlsx', *, sheets=1):
        p = self.folder / f'A사업장_20261028_{hour}{extension}'
        wb = Workbook()
        ws = wb.active
        ws.title = '원본시트'
        ws['A1'] = '00123'
        ws['A1'].font = Font(bold=True, color='00FF0000')
        ws['A1'].fill = PatternFill('solid', fgColor='00FFFF00')
        ws['B1'] = 42.5
        ws['B1'].number_format = '0.00'
        ws['C1'] = datetime(2026, 10, 28, 14, 30)
        ws['C1'].number_format = 'yyyy-mm-dd hh:mm'
        ws['D1'] = '=B1*2'
        ws['E1'] = '=SUM(A1:A2)'
        ws['E1'].data_type = 's'  # 원본의 수식처럼 보이는 텍스트
        ws['A2'] = '병합 제목'
        ws.merge_cells('A2:B2')
        ws.column_dimensions['B'].width = 24
        ws.row_dimensions[2].height = 30
        ws.freeze_panes = 'B2'
        if sheets > 1:
            wb.create_sheet('추가시트')
        wb.save(p)
        wb.close()
        return p

    def merge(self):
        return merge_group(scan_folder(self.folder).groups[0], self.out)

    def test_mixed_csv_xlsx_preserves_text_styles_types_and_merges(self):
        src_csv = self.csv('11시', 'ID,메모\n00123,=SUM(A1:A4)\n')
        src_xlsx = self.excel('14시')
        orig = [p.read_bytes() for p in (src_csv, src_xlsx)]
        result = self.merge()
        self.assertEqual(result.status, MergeStatus.CREATED)
        wb = load_workbook(result.path, data_only=False)
        try:
            self.assertEqual(wb.sheetnames, ['11시', '14시'])
            self.assertEqual(wb['11시']['A2'].value, '00123')
            self.assertEqual(wb['11시']['B2'].data_type, 's')
            self.assertEqual(wb['11시']['B2'].value, '=SUM(A1:A4)')
            sheet = wb['14시']
            self.assertEqual(sheet['A1'].value, '00123')
            self.assertTrue(sheet['A1'].font.bold)
            self.assertEqual(sheet['A1'].fill.fgColor.rgb, '00FFFF00')
            self.assertEqual(sheet['B1'].value, 42.5)
            self.assertEqual(sheet['B1'].number_format, '0.00')
            self.assertEqual(sheet['C1'].value, datetime(2026, 10, 28, 14, 30))
            self.assertEqual(sheet['D1'].value, '=B1*2')
            self.assertEqual(sheet['D1'].data_type, 'f')
            self.assertEqual(sheet['E1'].data_type, 's')
            self.assertEqual(str(next(iter(sheet.merged_cells.ranges))), 'A2:B2')
            self.assertEqual(sheet.column_dimensions['B'].width, 24)
            self.assertEqual(sheet.row_dimensions[2].height, 30)
            self.assertEqual(sheet.freeze_panes, 'B2')
        finally:
            wb.close()
        self.assertEqual([p.read_bytes() for p in (src_csv, src_xlsx)], orig)

    def test_excel_only_with_xlsm_and_xlsx(self):
        self.excel('11시', '.xlsx')
        self.excel('17시', '.xlsm')
        result = self.merge()
        w = load_workbook(result.path)
        try:
            self.assertEqual(w.sheetnames, ['11시', '17시'])
            self.assertEqual(w['17시']['B1'].value, 42.5)
        finally:
            w.close()

    def test_same_hour_csv_and_excel_create_distinct_sheet_names(self):
        self.csv('11시', '값\nCSV\n')
        self.excel('11-00')
        self.excel('17시')
        result = self.merge()
        w = load_workbook(result.path)
        try:
            self.assertEqual(w.sheetnames, ['11시', '11시 (2)', '17시'])
        finally:
            w.close()

    def test_multiple_worksheets_fail_without_partial_output(self):
        self.csv('11시', 'a\nb\n')
        self.excel('14시', sheets=2)
        with self.assertRaisesRegex(ValueError, '시트.*1개'):
            self.merge()
        self.assertEqual(list(self.out.glob('*')), [])

    def test_corrupt_xlsx_keeps_originals_and_removes_temporary_file(self):
        self.csv('11시', 'a\nb\n')
        damaged = self.folder / 'A사업장_20261028_14시.xlsx'
        damaged.write_bytes(b'BROKEN_FILE')
        with self.assertRaisesRegex(ValueError, "Excel 파일을 열 수 없습니다"):
            self.merge()
        self.assertEqual(damaged.read_bytes(), b'BROKEN_FILE')
        self.assertEqual(list(self.out.glob('*')), [])

    def test_existing_output_never_overwritten_on_mixed_group(self):
        self.csv('11시', 'a\nb\n')
        self.excel('14시')
        self.out.mkdir()
        path = self.out / 'A사업장_20261028.xlsx'
        path.write_bytes(b'EXISTING')
        self.assertEqual(self.merge().status, MergeStatus.EXISTS)
        self.assertEqual(path.read_bytes(), b'EXISTING')

    def test_legacy_xls_copies_string_number_boolean_and_date(self):
        self.csv('11시', 'ID\n0001\n')
        legacy = self.folder / 'A사업장_20261028_14시.xls'
        legacy.write_bytes(b'FAKE-LEGACY')
        fake = types.ModuleType('xlrd')
        fake.XL_CELL_EMPTY = 0
        fake.XL_CELL_TEXT = 1
        fake.XL_CELL_NUMBER = 2
        fake.XL_CELL_DATE = 3
        fake.XL_CELL_BOOLEAN = 4
        fake.XL_CELL_ERROR = 5
        fake.XL_CELL_BLANK = 6
        fake.error_text_from_code = {7: '#DIV/0!'}
        fake.xldate_as_datetime = lambda value, mode: datetime(2026, 10, 28)
        class LegacySheet:
            nrows, ncols = 2, 4
            _rows = [[(1,'00123'), (2,9.5), (4,1), (3,42)],
                     [(1,'=SUM(A1)'), (2,12), (0,''), (5,7)]]
            def cell(self, r, c):
                typ, val = self._rows[r][c]
                return types.SimpleNamespace(ctype=typ, value=val)
        class LegacyWorkbook:
            nsheets, datemode = 1, 0
            def sheet_by_index(self, n): return LegacySheet()
            def release_resources(self): pass
        fake.open_workbook = lambda path, on_demand: LegacyWorkbook()
        with patch.dict(sys.modules, {'xlrd': fake}):
            result = self.merge()
        w = load_workbook(result.path)
        try:
            self.assertEqual(w.sheetnames, ['11시', '14시'])
            ws = w['14시']
            self.assertEqual(ws['A1'].value, '00123')
            self.assertEqual(ws['A2'].value, '=SUM(A1)')
            self.assertEqual(ws['A2'].data_type, 's')
            self.assertEqual(ws['B1'].value, 9.5)
            self.assertIs(ws['C1'].value, True)
            self.assertEqual(ws['D1'].value.date().isoformat(), '2026-10-28')
            self.assertEqual(ws['D2'].data_type, 'e')
        finally:
            w.close()

if __name__ == '__main__':
    unittest.main()
