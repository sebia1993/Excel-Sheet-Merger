"""Handled merge failures release only their own XML and OS handles.

Temporary-file registry inspection is test-only. Production must never use a
registry snapshot to decide which other workbooks' files it may remove.
"""
from __future__ import annotations

import errno
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet._writer import ALL_TEMP_FILES, WorksheetWriter
from excel_sheet_merger.classify import scan_folder
from excel_sheet_merger.merge import _append_openpyxl, merge_group


class WorkbookResourceCleanupTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.output = self.root / '병합결과'
        self.previous_registry = set(ALL_TEMP_FILES)
        self.retained_archives = []
        self.retained_errors = []
        self.owned_books = []

    def tearDown(self):
        # Repair the intentionally failing baseline experiments, not production.
        for archive in self.retained_archives:
            archive.close()
        for book in self.owned_books:
            for sheet in book.worksheets:
                if book.write_only:
                    try:
                        if not sheet.closed:
                            sheet.close()
                    except Exception:
                        pass
            book.close()
        for name in set(ALL_TEMP_FILES) - self.previous_registry:
            try:
                Path(name).unlink(missing_ok=True)
            finally:
                if name in ALL_TEMP_FILES:
                    ALL_TEMP_FILES.remove(name)
        self.temp.cleanup()

    def csv(self, hour, body='ID,주소\n00123,192.0.2.1\n'):
        p = self.root / f'A사업장_20261028_{hour}.csv'
        p.write_text(body, encoding='utf-8')
        return p

    def xlsx(self, hour):
        p = self.root / f'A사업장_20261028_{hour}.xlsx'
        b = Workbook()
        b.active['A1'] = 'Excel'
        b.active['B1'] = '=1+2'
        b.active['B1'].font = Font(bold=True)
        b.save(p)
        b.close()
        return p

    def assert_no_owned_temp(self, before):
        extra = set(ALL_TEMP_FILES) - before
        self.assertEqual(extra, set(), f'owned worksheet XML still registered: {extra}')
        self.assertFalse(any(self.output.glob('_병합임시_*')))

    def fail_merge(self, expected=None):
        try:
            merge_group(scan_folder(self.root).groups[0], self.output)
        except Exception as exc:
            self.retained_errors.append(exc)
            if expected is not None:
                self.assertIn(expected, str(exc))
            return exc
        self.fail('fixture must fail, not publish a workbook')

    def test_empty_second_csv_reclaims_first_streamed_xml(self):
        self.csv('11시')
        self.csv('14시', '')
        before = set(ALL_TEMP_FILES)
        self.fail_merge('비어')
        self.assert_no_owned_temp(before)
        self.assertEqual(list(self.output.glob('*.xlsx')), [])

    def test_corrupt_excel_after_csv_reclaims_streamed_xml(self):
        self.csv('11시')
        (self.root / 'A사업장_20261028_14시.xlsx').write_bytes(b'not an XLSX')
        before = set(ALL_TEMP_FILES)
        self.fail_merge('Excel')
        self.assert_no_owned_temp(before)

    def test_invalid_csv_after_written_rows_reclaims_xml(self):
        self.csv('11시', 'ID\n001\ninvalid\x01\n')
        self.csv('14시')
        before = set(ALL_TEMP_FILES)
        self.fail_merge('제어 문자')
        self.assert_no_owned_temp(before)

    def _save_failure(self, csv_only):
        if csv_only:
            self.csv('11시')
            self.csv('14시')
        else:
            self.xlsx('11시')
            self.xlsx('14시')
        self.output.mkdir()
        previous = self.output / 'A사업장_20261028.xlsx'
        previous.write_bytes(b'KEEP EXISTING')
        before = set(ALL_TEMP_FILES)
        original = zipfile.ZipFile.write
        failure = OSError(errno.ENOSPC, 'injected archive disk-full')

        def fail_sheet(archive, filename, arcname=None, *args, **kwargs):
            if arcname and 'worksheets/sheet' in str(arcname):
                self.retained_archives.append(archive)
                raise failure
            return original(archive, filename, arcname, *args, **kwargs)

        with patch.object(zipfile.ZipFile, 'write', fail_sheet):
            error = self.fail_merge('injected archive disk-full')
        self.assertIs(error, failure, 'cleanup must not replace the original failure')
        self.assertEqual(previous.read_bytes(), b'KEEP EXISTING')
        self.assertEqual(list(self.output.glob('*.xlsx')), [previous])
        for archive in self.retained_archives:
            self.assertTrue(archive.fp is None or archive.fp.closed,
                            'output archive handle retained after handled failure')
        self.assert_no_owned_temp(before)

    def test_normal_excel_save_failure_closes_zip_and_cleans_xml(self):
        self._save_failure(csv_only=False)

    def test_streamed_csv_save_failure_closes_zip_and_cleans_xml(self):
        self._save_failure(csv_only=True)

    def test_malformed_source_closes_input_even_with_retained_traceback(self):
        source = self.xlsx('11시')
        broken = self.root / 'broken.xlsx'
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(broken, 'w') as target:
            for info in original.infolist():
                target.writestr(info, b'<invalid XML' if info.filename == 'xl/workbook.xml'
                                else original.read(info))
        real = zipfile.ZipFile
        def record(*args, **kwargs):
            archive = real(*args, **kwargs)
            self.retained_archives.append(archive)
            return archive
        output_book = Workbook()
        self.owned_books.append(output_book)
        with patch('openpyxl.reader.excel.ZipFile', side_effect=record):
            try:
                _append_openpyxl(output_book.active, broken)
            except ValueError as exc:
                self.retained_errors.append(exc)
                self.assertIn('broken.xlsx', str(exc))
            else:
                self.fail('malformed XML must be rejected')
        self.assertTrue(self.retained_archives)
        for archive in self.retained_archives:
            self.assertTrue(archive.fp is None or archive.fp.closed,
                            'input stream retained while traceback is live')

    def test_other_live_writer_survives_failed_merge(self):
        foreign = Workbook(write_only=True)
        self.owned_books.append(foreign)
        sheet = foreign.create_sheet('other task')
        sheet.append(['do not delete'])
        unrelated = Path(sheet._writer.out)
        before = set(ALL_TEMP_FILES)
        self.csv('11시')
        self.csv('14시', '')
        self.fail_merge('비어')
        self.assertTrue(unrelated.exists(), 'cleanup touched an unrelated active writer')
        self.assert_no_owned_temp(before)
        sheet.append(['still usable'])
        result = self.root / 'other.xlsx'
        foreign.save(result)
        with zipfile.ZipFile(result) as package:
            self.assertIsNone(package.testzip())
        opened = load_workbook(result)
        try:
            self.assertEqual(opened.active['A2'].value, 'still usable')
        finally:
            opened.close()

    def test_repeated_failures_then_success_leave_no_owned_xml(self):
        self.csv('11시')
        second = self.csv('14시', '')
        before = set(ALL_TEMP_FILES)
        for _ in range(12):
            self.fail_merge('비어')
            self.assert_no_owned_temp(before)
        second.write_text('ID\n00456\n', encoding='utf-8')
        result = merge_group(scan_folder(self.root).groups[0], self.output)
        self.assert_no_owned_temp(before)
        opened = load_workbook(result.path)
        try:
            self.assertEqual(opened.sheetnames, ['11시', '14시'])
            self.assertEqual(opened['11시']['A2'].value, '00123')
        finally:
            opened.close()

    def test_safe_serialization_matches_standard_save_package(self):
        from xml.etree import ElementTree as ET
        from openpyxl.comments import Comment
        from excel_sheet_merger.workbook_resources import save_workbook_safely
        def build():
            book = Workbook()
            sheet = book.active
            sheet.title = '11시'
            sheet['A1'] = '00123'
            sheet['B1'] = '=SUM(1,2)'
            sheet['B1'].font = Font(bold=True, color='166D53')
            sheet['A2'] = '메모'
            sheet['A2'].comment = Comment('original comment', 'test')
            sheet['B2'] = 'link'
            sheet['B2'].hyperlink = 'https://example.invalid/'
            sheet['C4'] = 'merged'
            sheet.merge_cells('C4:D4')
            sheet.row_dimensions[4].height = 31
            sheet.column_dimensions['B'].width = 24
            sheet.freeze_panes = 'B2'
            sheet.auto_filter.ref = 'A1:B2'
            return book
        reference, candidate = build(), build()
        self.owned_books.extend([reference, candidate])
        a, b = self.root/'reference.xlsx', self.root/'candidate.xlsx'
        reference.save(a)
        before = set(ALL_TEMP_FILES)
        save_workbook_safely(candidate, b)
        self.assert_no_owned_temp(before)
        with zipfile.ZipFile(a) as left, zipfile.ZipFile(b) as right:
            self.assertIsNone(right.testzip())
            self.assertEqual(left.namelist(), right.namelist())
            for name in left.namelist():
                ldata, rdata = left.read(name), right.read(name)
                if name == 'docProps/core.xml':
                    def without_modified(payload):
                        element = ET.fromstring(payload)
                        for item in list(element):
                            if item.tag.endswith('}modified'):
                                element.remove(item)
                        return ET.tostring(element)
                    ldata, rdata = without_modified(ldata), without_modified(rdata)
                self.assertEqual(ldata, rdata, name)

    def test_mixed_success_and_numbered_rerun_preserve_values(self):
        self.csv('11시', 'ID,메모\n00123,=SUM(1)\n')
        source = self.xlsx('14시')
        original = source.read_bytes()
        before = set(ALL_TEMP_FILES)
        group = scan_folder(self.root).groups[0]
        first = merge_group(group, self.output)
        first_bytes = first.path.read_bytes()
        second = merge_group(group, self.output)
        self.assertEqual(second.path.name, 'A사업장_20261028 (2).xlsx')
        self.assertEqual(first.path.read_bytes(), first_bytes)
        self.assertEqual(source.read_bytes(), original)
        self.assert_no_owned_temp(before)
        result = load_workbook(second.path)
        try:
            self.assertEqual(result.sheetnames, ['11시', '14시'])
            self.assertEqual(result['11시']['A2'].value, '00123')
            self.assertEqual(result['11시']['B2'].data_type, 's')
            self.assertEqual(result['14시']['B1'].value, '=1+2')
            self.assertTrue(result['14시']['B1'].font.bold)
        finally:
            result.close()

    def test_secondary_cleanup_error_does_not_hide_source_failure(self):
        self.csv('11시')
        self.csv('14시', '')
        with patch.object(WorksheetWriter, 'cleanup', side_effect=PermissionError('locked temp')):
            self.fail_merge('비어')
        self.assertEqual(list(self.output.glob('*.xlsx')), [])


if __name__ == '__main__':
    unittest.main()
