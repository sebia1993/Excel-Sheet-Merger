"""Low disk-space safety: fail before final ZIP, not after publishing results."""
from __future__ import annotations

import errno
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from excel_sheet_merger.classify import scan_folder
from excel_sheet_merger.merge import merge_group


class DiskCapacityTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        (self.root / 'A사업장_20261028_11시.csv').write_text(
            'ID,메모\n00123,=SUM(A1:A3)\n', encoding='utf-8')
        wb = Workbook()
        wb.active['A1'] = '엑셀'
        wb.active['B2'] = '=1+2'
        wb.save(self.root / 'A사업장_20261028_14시.xlsx')
        wb.close()
        self.group = scan_folder(self.root).groups[0]
        self.dest = self.root / '병합결과'

    def test_low_free_space_stops_before_final_stitch_and_keeps_previous_output(self):
        self.dest.mkdir()
        previous = self.dest / 'A사업장_20261028.xlsx'
        previous.write_bytes(b'PREVIOUS-NEVER-OVERWRITE')
        no_space = shutil._ntuple_diskusage(10 * 1024**3, 10 * 1024**3 - 1, 1)
        with patch('excel_sheet_merger.merge.shutil.disk_usage', return_value=no_space), \
             patch('excel_sheet_merger.merge._stitch_csv_sheets', side_effect=AssertionError('preflight was skipped')) as stitch:
            with self.assertRaises(OSError) as caught:
                merge_group(self.group, self.dest)
        self.assertEqual(caught.exception.errno, errno.ENOSPC)
        self.assertIn('공간', str(caught.exception))
        stitch.assert_not_called()
        self.assertEqual(previous.read_bytes(), b'PREVIOUS-NEVER-OVERWRITE')
        self.assertEqual(list(self.dest.glob('_병합임시_*')), [])
        self.assertEqual(list(self.dest.glob('*.xlsx')), [previous])

    def test_missing_disk_free_space_metadata_does_not_block_valid_mixed_merge(self):
        with patch('excel_sheet_merger.merge.shutil.disk_usage', side_effect=OSError('SMB free space unavailable')):
            result = merge_group(self.group, self.dest)
        wb = load_workbook(result.path)
        try:
            self.assertEqual(wb.sheetnames, ['11시', '14시'])
            self.assertEqual(wb['11시']['A2'].value, '00123')
            self.assertEqual(wb['14시']['B2'].value, '=1+2')
        finally:
            wb.close()
        self.assertEqual(list(self.dest.glob('_병합임시_*')), [])

    def test_sufficient_disk_space_keeps_mixed_workbook_semantics(self):
        plentiful = shutil._ntuple_diskusage(100 * 1024**3, 1 * 1024**3, 99 * 1024**3)
        with patch('excel_sheet_merger.merge.shutil.disk_usage', return_value=plentiful) as disk_query:
            outcome = merge_group(self.group, self.dest)
        disk_query.assert_called_once_with(self.dest)
        wb = load_workbook(outcome.path)
        try:
            self.assertEqual(wb['11시']['B2'].value, '=SUM(A1:A3)')
            self.assertEqual(wb['11시']['B2'].data_type, 's')
            self.assertEqual(wb['14시']['B2'].value, '=1+2')
            self.assertEqual(wb['14시']['B2'].data_type, 'f')
        finally:
            wb.close()

    def test_disk_fills_after_check_never_publishes_partial_output(self):
        self.dest.mkdir()
        previous = self.dest / 'A사업장_20261028.xlsx'
        previous.write_bytes(b'PREVIOUS')
        plentiful = shutil._ntuple_diskusage(100 * 1024**3, 1 * 1024**3, 99 * 1024**3)
        with patch('excel_sheet_merger.merge.shutil.disk_usage', return_value=plentiful), \
             patch('excel_sheet_merger.merge._stitch_csv_sheets', side_effect=OSError(errno.ENOSPC, 'No space left on device')):
            with self.assertRaises(OSError) as caught:
                merge_group(self.group, self.dest)
        self.assertEqual(caught.exception.errno, errno.ENOSPC)
        self.assertEqual(previous.read_bytes(), b'PREVIOUS')
        self.assertEqual(list(self.dest.glob('*.xlsx')), [previous])
        self.assertFalse(list(self.dest.glob('_병합임시_*')))

    def test_csv_only_is_not_subject_to_mixed_zip_disk_preflight(self):
        (self.root / 'A사업장_20261028_14시.xlsx').unlink()
        (self.root / 'A사업장_20261028_14시.csv').write_text('ID\n00999\n', encoding='utf-8')
        group = scan_folder(self.root).groups[0]
        with patch('excel_sheet_merger.merge.shutil.disk_usage', side_effect=AssertionError('CSV-only path should remain unchanged')):
            outcome = merge_group(group, self.dest)
        wb = load_workbook(outcome.path)
        try:
            self.assertEqual(wb.sheetnames, ['11시', '14시'])
            self.assertEqual(wb['14시']['A2'].value, '00999')
        finally:
            wb.close()


if __name__ == '__main__':
    unittest.main()
