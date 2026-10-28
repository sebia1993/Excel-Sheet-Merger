import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from excel_sheet_merger.classify import parse_filename, scan_folder, unique_sheet_names


class ClassifyTests(unittest.TestCase):
    def test_example_three_files_same_group_sorted_by_time(self):
        with tempfile.TemporaryDirectory() as dirname:
            p = Path(dirname)
            for filename in ['A사업장_20261028_17시.csv', 'A사업장_20261028_11시.csv', 'A사업장_20261028_14시.csv']:
                (p / filename).touch()
            result = scan_folder(p)
            self.assertEqual(len(result.groups), 1)
            self.assertEqual(result.groups[0].output_name, 'A사업장_20261028.xlsx')
            self.assertEqual([x.sheet_name for x in result.groups[0].inputs], ['11시', '14시', '17시'])

    def test_supported_date_and_time_formats_merge(self):
        names = ['서울R&D_2026-10-28_11시.csv','서울R&D_2026_10_28_14-30.csv', '서울R&D_2026.10.28_1700.csv']
        parsed = [parse_filename(Path(s)) for s in names]
        self.assertEqual({item.site for item in parsed}, {'서울R&D'})
        self.assertEqual({item.date for item in parsed}, {'20261028'})
        self.assertEqual([(i.hour,i.minute) for i in parsed], [(11,0),(14,30),(17,0)])

    def test_different_business_is_never_fuzzy_merged(self):
        with tempfile.TemporaryDirectory() as d:
            for s in ['A사업장_본관_20261028_11시.csv','A사업장_별관_20261028_14시.csv']:
                (Path(d) / s).touch()
            result = scan_folder(Path(d))
            self.assertEqual(len(result.groups), 0)
            self.assertEqual(len(result.singles), 2)

    def test_trailing_qualifier_not_discarded(self):
        a = parse_filename(Path('A사업장_20261028_11시_최종.csv'))
        b = parse_filename(Path('A사업장_20261028_14시.csv'))
        self.assertNotEqual(a.site, b.site)

    def test_duplicate_sheet_names(self):
        self.assertEqual(unique_sheet_names(['11시','11시','14시','11시']), ['11시', '11시 (2)', '14시', '11시 (3)'])

    def test_ambiguous_or_invalid_names_are_excluded(self):
        invalid = ['A사업장_20261328_11시.csv','A사업장_20261028_11시_14시.csv', 'A사업장_20261028.csv','A사업장_20261028_25시.csv']
        with tempfile.TemporaryDirectory() as d:
            for name in invalid:
                (Path(d) / name).touch()
            result = scan_folder(Path(d))
            self.assertEqual(len(result.groups), 0)
            self.assertEqual(len(result.rejected), len(invalid))

    def test_csv_and_excel_suffixes_case_insensitive(self):
        for ext in (".csv", ".CSV", ".xlsx", ".XLSX", ".xlsm", ".xls"):
            with self.subTest(ext=ext):
                self.assertEqual(parse_filename(Path(f"A사업장_20261028_11시{ext}")).sheet_name, "11시")
        with self.assertRaisesRegex(ValueError, "지원하지"):
            parse_filename(Path("A사업장_20261028_11시.xlsb"))

    def test_mixed_csv_and_excel_are_grouped_and_sorted_by_time(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder)
            for filename in ("A사업장_20261028_17시.xlsx", "A사업장_20261028_11시.csv",
                             "A사업장_20261028_14시.xls", "B사업장_20261028_11시.xlsm"):
                (source / filename).touch()
            result = scan_folder(source)
            self.assertEqual(len(result.groups), 1)
            self.assertEqual([x.sheet_name for x in result.groups[0].inputs], ["11시", "14시", "17시"])
            self.assertEqual([x.path.suffix for x in result.groups[0].inputs], [".csv", ".xls", ".xlsx"])
            self.assertEqual(len(result.singles), 1)
            self.assertEqual(len(result.rejected), 0)

    def test_long_site_names_keep_dates_and_publish_distinct_results(self):
        from openpyxl import load_workbook
        from excel_sheet_merger.merge import merge_group

        long_site = 'Office' + 'A' * 160
        with tempfile.TemporaryDirectory() as name:
            folder = Path(name)
            for date in ('20261028', '20261029'):
                for hour in ('11시', '14시'):
                    (folder / f'{long_site}_{date}_{hour}.csv').write_text(
                        f'date,ID\n{date},00123\n', encoding='utf8')
            result = scan_folder(folder)
            self.assertEqual(len(result.groups), 2)
            outputs = {g.date: g.output_name for g in result.groups}
            for date in ('20261028','20261029'):
                self.assertTrue(outputs[date].endswith(f'_{date}.xlsx'))
                self.assertLessEqual(len(outputs[date]), 155)
            self.assertNotEqual(outputs['20261028'], outputs['20261029'])
            for g in result.groups:
                merged = merge_group(g, folder / '병합결과')
                self.assertEqual(merged.path.name, outputs[g.date])
                wb = load_workbook(merged.path, read_only=True)
                try:
                    self.assertEqual(wb.sheetnames, ['11시','14시'])
                    self.assertEqual(wb['11시']['A2'].value, g.date)
                    self.assertEqual(wb['11시']['B2'].value, '00123')
                finally:
                    wb.close()
            self.assertEqual(len(list((folder/'병합결과').glob('*.xlsx'))), 2)

    def test_long_sites_with_same_truncated_prefix_stay_separate(self):
        site_a = 'Office' + 'B' * 160 + 'X'
        site_b = 'Office' + 'B' * 160 + 'Y'
        with tempfile.TemporaryDirectory() as name:
            folder = Path(name)
            for site in (site_a, site_b):
                for hour in ('11시', '14시'):
                    (folder / f'{site}_20261029_{hour}.csv').write_text(
                        'ID\n00123\n', encoding='utf8')
            report = scan_folder(folder)
            self.assertEqual(len(report.groups), 2)
            names = [g.output_name for g in report.groups]
            self.assertEqual(len(set(names)), 2)
            self.assertTrue(all(v.endswith('_20261029.xlsx') for v in names))
            self.assertTrue(all(len(v) <= 155 for v in names))
            self.assertEqual({g.site for g in report.groups}, {site_a, site_b})

    def test_12_hour_period_conversion(self):
        self.assertEqual((parse_filename(Path('사업장_20261028_오후2시.csv')).hour, parse_filename(Path('사업장_20261028_오전12시.csv')).hour), (14, 0))
        self.assertEqual(parse_filename(Path('사업장_20261028_오후12시.csv')).hour, 12)

    def test_multiple_formats_order_and_locks(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d)
            for s in ['~$A사업장_20261028_11시.csv','A사업장_20261028_오후2시.csv','20261028_A사업장_09시30분.csv','A사업장_20261028_15시.xlsm']:
                (p / s).touch()
            r = scan_folder(p)
            self.assertEqual([i.sheet_name for i in r.groups[0].inputs], ['09시30분', '14시', '15시'])
            self.assertEqual(len(r.rejected), 1)

    def test_excel_sheet_name_utf16_and_casefold(self):
        names = unique_sheet_names(['a'*40,'a'*40,'B'*30,'b'*30,'시트/1'])
        self.assertEqual(names[1][-4:], ' (2)')
        self.assertEqual(names[3][-4:], ' (2)')
        self.assertEqual(names[4], '시트-1')
        self.assertTrue(all(len(s.encode('utf-16-le')) // 2 <= 31 for s in names))


if __name__ == '__main__':
    unittest.main()
