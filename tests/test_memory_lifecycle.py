"""Large source cells should stop consuming memory while their results are copied.

These tests exercise real openpyxl workbooks, not a fake data model; observing
source-cell ownership during transfer catches a future regression to the
old high-peak-memory algorithm.
"""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from excel_sheet_merger.merge import _append_openpyxl


class ExcelSourceMemoryLifecycleTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.source = self.folder / 'A사업장_20261028_11시.xlsx'

    def test_live_source_cell_ownership_decreases_during_copy(self):
        book = Workbook()
        for i in range(1, 401):
            for j in range(1, 11):
                book.active.cell(i, j, f'업무_{i}_{j}')
        book.save(self.source)
        book.close()

        loaded = []
        target_book = Workbook()
        target = target_book.active
        original_cell = target.cell
        calls = 0
        live_at_midpoint = None

        def load_source(*args, **kwargs):
            source = load_workbook(*args, **kwargs)
            loaded.append(source)
            return source

        def tracking_copy(*args, **kwargs):
            nonlocal calls, live_at_midpoint
            calls += 1
            if calls == 2000:
                live_at_midpoint = len(loaded[0].active._cells)
            return original_cell(*args, **kwargs)

        with patch('excel_sheet_merger.merge.load_workbook', side_effect=load_source):
            with patch.object(target, 'cell', side_effect=tracking_copy):
                _append_openpyxl(target, self.source)
        self.assertEqual(calls, 4000)
        self.assertIsNotNone(live_at_midpoint)
        self.assertLessEqual(live_at_midpoint, 2500,
                             'source cells must be released before destination copy finishes')
        self.assertEqual(target['J400'].value, '업무_400_10')
        target_book.close()

    def test_styles_comments_merges_and_sparse_cells_survive_copy(self):
        book = Workbook()
        ws = book.active
        ws['B2'] = '=1+2'
        ws['B2'].font = Font(bold=True, color='008000')
        ws['C3'] = '문서 링크'
        ws['C3'].hyperlink = 'https://example.com/docs'
        ws['C3'].comment = Comment('관리 메모', 'QA')
        ws['Z200'] = '마지막 값'
        ws['D8'] = '헤더'
        ws.merge_cells('D8:F8')
        ws['H10'].fill = PatternFill(fill_type='solid', fgColor='FF00CC')
        ws.column_dimensions['B'].width = 32
        ws.row_dimensions[8].height = 44
        ws.freeze_panes = 'C5'
        ws.auto_filter.ref = 'B2:C3'
        book.save(self.source)
        book.close()
        saved_source = self.source.read_bytes()

        target_book = Workbook()
        _append_openpyxl(target_book.active, self.source)
        output = self.folder / 'copied.xlsx'
        target_book.save(output)
        target_book.close()
        copied = load_workbook(output)
        try:
            sheet = copied.active
            self.assertEqual(sheet['B2'].value, '=1+2')
            self.assertTrue(sheet['B2'].font.bold)
            self.assertEqual(sheet['C3'].hyperlink.target, 'https://example.com/docs')
            self.assertEqual(sheet['C3'].comment.text, '관리 메모')
            self.assertEqual(sheet['Z200'].value, '마지막 값')
            self.assertIn('D8:F8', {str(r) for r in sheet.merged_cells.ranges})
            self.assertEqual(sheet['H10'].fill.fgColor.rgb, '00FF00CC')
            self.assertEqual(sheet.column_dimensions['B'].width, 32)
            self.assertEqual(sheet.row_dimensions[8].height, 44)
            self.assertEqual(sheet.freeze_panes, 'C5')
            self.assertEqual(sheet.auto_filter.ref, 'B2:C3')
        finally:
            copied.close()
        self.assertEqual(self.source.read_bytes(), saved_source,
                         'the on-disk source must remain unmodified')


if __name__ == '__main__':
    unittest.main()
