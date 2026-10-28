"""Source ZIP resource safety is checked before openpyxl inflates workbook parts."""
from __future__ import annotations

import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from excel_sheet_merger.classify import scan_folder
from excel_sheet_merger.merge import SourceFileError, merge_group


class ArchiveResourceSafetyTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.output = self.folder / '병합결과'
        (self.folder / 'A사업장_20261028_11시.csv').write_text('ID\n00123\n', encoding='utf-8')
        self.excel_path = self.folder / 'A사업장_20261028_14시.xlsx'
        wb = Workbook()
        wb.active['A1'] = '정상 데이터'
        wb.active['B1'] = '=1+2'
        wb.save(self.excel_path)
        wb.close()

    def _replace_member(self, member_name, payload):
        rewritten = self.folder / 'altered.xlsx'
        with zipfile.ZipFile(self.excel_path, 'r') as source, \
                zipfile.ZipFile(rewritten, 'w') as result:
            for entry in source.infolist():
                body = payload if entry.filename == member_name else source.read(entry.filename)
                result.writestr(entry, body)
        rewritten.replace(self.excel_path)

    def _merge(self):
        return merge_group(scan_folder(self.folder).groups[0], self.output)

    def test_pathological_valid_xml_expansion_is_rejected_before_parsing(self):
        # A real ZIP with a well-formed sheet that expands to >64MiB, from
        # much less than 1MiB compressed. Repeated short comments avoid long
        # XML text-node parser limits and are harmless to Excel cell values.
        comments = b'<!--' + b'A' * (32 * 1024) + b'-->'
        with zipfile.ZipFile(self.excel_path, 'r') as book:
            original = book.read('xl/worksheets/sheet1.xml')
        payload = original.replace(b'</worksheet>', comments * 2110 + b'</worksheet>')
        self.assertGreater(len(payload), 64 * 1024 * 1024)
        self._replace_member('xl/worksheets/sheet1.xml', payload)
        self.assertLess(self.excel_path.stat().st_size, 512 * 1024)
        preserved = self.excel_path.read_bytes()
        with patch('excel_sheet_merger.merge.load_workbook', wraps=load_workbook) as parse:
            with self.assertRaisesRegex(SourceFileError, '압축 해제|비정상적으로 큰'):
                self._merge()
            parse.assert_not_called()
        self.assertEqual(self.excel_path.read_bytes(), preserved)
        self.assertEqual(list(self.output.iterdir()), [])

    def test_unread_custom_xml_part_with_high_ratio_does_not_block_real_sheet(self):
        # Office Custom XML auxiliary parts are not read by pinned openpyxl.
        # This tests a real high-compression ZIP, not just a mocked ZipInfo.
        document = b'<root>' + b'A' * 160_000 + b'</root>'
        with zipfile.ZipFile(self.excel_path, 'a') as book:
            book.writestr('customXml/item1.xml', document, compress_type=zipfile.ZIP_DEFLATED)
        with zipfile.ZipFile(self.excel_path) as book:
            member = book.getinfo('customXml/item1.xml')
            self.assertGreater(member.file_size, 100_000)
            self.assertGreater(member.file_size, member.compress_size * 400)
        source_before = self.excel_path.read_bytes()
        with patch('excel_sheet_merger.xlsx_safety.SUSPICIOUS_EXPANDED_MEMBER_BYTES', 8 * 1024):
            result = self._merge()
        self.assertEqual(self.excel_path.read_bytes(), source_before)
        book = load_workbook(result.path)
        try:
            self.assertEqual(book.sheetnames, ['11시', '14시'])
            self.assertEqual(book['11시']['A2'].value, '00123')
            self.assertEqual(book['14시']['A1'].value, '정상 데이터')
            self.assertEqual(book['14시']['B1'].value, '=1+2')
        finally:
            book.close()

    def test_high_ratio_actual_worksheet_is_still_blocked_before_parser(self):
        comments = b'<!--' + b'Q' * 320_000 + b'-->'
        with zipfile.ZipFile(self.excel_path) as book:
            original = book.read('xl/worksheets/sheet1.xml')
        self._replace_member('xl/worksheets/sheet1.xml',
                             original.replace(b'</worksheet>', comments + b'</worksheet>'))
        with patch('excel_sheet_merger.xlsx_safety.SUSPICIOUS_EXPANDED_MEMBER_BYTES', 8 * 1024):
            with patch('excel_sheet_merger.merge.load_workbook', wraps=load_workbook) as parser:
                with self.assertRaisesRegex(SourceFileError, '압축 확장 비율'):
                    self._merge()
                parser.assert_not_called()
        self.assertEqual(list(self.output.iterdir()), [])

    def test_nonstandard_customxml_path_still_enforces_high_ratio(self):
        with zipfile.ZipFile(self.excel_path, 'a') as book:
            book.writestr('customXml/other.xml', b'<root>' + b'B' * 160_000 + b'</root>',
                          compress_type=zipfile.ZIP_DEFLATED)
        with patch('excel_sheet_merger.xlsx_safety.SUSPICIOUS_EXPANDED_MEMBER_BYTES', 8 * 1024):
            with patch('excel_sheet_merger.merge.load_workbook', wraps=load_workbook) as parser:
                with self.assertRaisesRegex(SourceFileError, '압축 확장 비율'):
                    self._merge()
                parser.assert_not_called()

    def test_ignored_customxml_is_still_subject_to_expanded_xml_size_cap(self):
        with zipfile.ZipFile(self.excel_path, 'a') as book:
            book.writestr('customXml/item1.xml', b'<root>' + b'C' * 100_000 + b'</root>',
                          compress_type=zipfile.ZIP_DEFLATED)
        with patch('excel_sheet_merger.xlsx_safety.MAX_EXPANDED_XML_BYTES', 10 * 1024):
            with patch('excel_sheet_merger.merge.load_workbook', wraps=load_workbook) as parser:
                with self.assertRaisesRegex(SourceFileError, 'XML 용량 한도'):
                    self._merge()
                parser.assert_not_called()

    def test_duplicate_excel_part_cannot_be_loaded_ambiguously(self):
        with zipfile.ZipFile(self.excel_path) as book:
            sheet_xml = book.read('xl/worksheets/sheet1.xml')
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter('ignore', UserWarning)
            with zipfile.ZipFile(self.excel_path, 'a') as book:
                book.writestr('xl/worksheets/sheet1.xml', sheet_xml)
        preserved = self.excel_path.read_bytes()
        with self.assertRaisesRegex(SourceFileError, '중복된 구성 요소'):
            self._merge()
        self.assertEqual(self.excel_path.read_bytes(), preserved)
        self.assertEqual(list(self.output.iterdir()), [])

    def test_duplicate_package_manifest_is_rejected_before_loading(self):
        # Not all ambiguous package parts live below xl/; a duplicated
        # content-type manifest can also change interpretation by ZIP reader.
        with zipfile.ZipFile(self.excel_path) as book:
            manifest = book.read('[Content_Types].xml')
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter('ignore', UserWarning)
            with zipfile.ZipFile(self.excel_path, 'a') as book:
                book.writestr('[Content_Types].xml', manifest)
        with patch('excel_sheet_merger.merge.load_workbook', wraps=load_workbook) as parser:
            with self.assertRaisesRegex(SourceFileError, '중복된 구성 요소'):
                self._merge()
            parser.assert_not_called()
        self.assertEqual(list(self.output.iterdir()), [])

    def test_excessive_cumulative_zip_metadata_stops_before_parsing(self):
        # A lowered test-only safety budget checks the aggregate condition
        # without spending gigabytes of RAM or disk on a CI runner.
        with patch('excel_sheet_merger.xlsx_safety.MAX_EXPANDED_PACKAGE_BYTES', 400):
            with patch('excel_sheet_merger.merge.load_workbook', wraps=load_workbook) as parser:
                with self.assertRaisesRegex(SourceFileError, '전체 용량 한도'):
                    self._merge()
                parser.assert_not_called()
        self.assertEqual(list(self.output.iterdir()), [])

    def test_xml_member_budget_blocks_expansion_without_parsing(self):
        with patch('excel_sheet_merger.xlsx_safety.MAX_EXPANDED_XML_BYTES', 500):
            with patch('excel_sheet_merger.merge.load_workbook', wraps=load_workbook) as parser:
                with self.assertRaisesRegex(SourceFileError, 'XML 용량 한도'):
                    self._merge()
                parser.assert_not_called()
        self.assertEqual(list(self.output.iterdir()), [])

    def test_members_limit_prevents_pathological_archive(self):
        with patch('excel_sheet_merger.xlsx_safety.MAX_ARCHIVE_MEMBERS', 4):
            with patch('excel_sheet_merger.merge.load_workbook', wraps=load_workbook) as parser:
                with self.assertRaisesRegex(SourceFileError, '많은 구성 요소'):
                    self._merge()
                parser.assert_not_called()
        self.assertEqual(list(self.output.iterdir()), [])

    def test_existing_result_is_preserved_when_next_excel_is_unsafe(self):
        self.output.mkdir()
        previous = self.output / 'A사업장_20261028.xlsx'
        previous.write_bytes(b'PREEXISTING')
        with patch('excel_sheet_merger.xlsx_safety.MAX_EXPANDED_PACKAGE_BYTES', 400):
            with self.assertRaises(SourceFileError):
                self._merge()
        self.assertEqual(previous.read_bytes(), b'PREEXISTING')
        self.assertEqual([p.name for p in self.output.iterdir()], [previous.name])

    def test_resource_limit_has_specific_actionable_korean_guidance(self):
        from excel_sheet_merger.failure_guidance import diagnose_failure
        with patch('excel_sheet_merger.xlsx_safety.MAX_EXPANDED_PACKAGE_BYTES', 400):
            with self.assertRaises(SourceFileError) as failure:
                self._merge()
        advice = diagnose_failure(failure.exception)
        self.assertEqual(advice.filename, self.excel_path.name)
        self.assertEqual(advice.code, 'EXCEL_RESOURCE_LIMIT')
        self.assertIn('압축', advice.reason)
        self.assertIn('확인', advice.action)
        self.assertNotIn('암호화되었', advice.reason)
        self.assertEqual(list(self.output.iterdir()), [])

    def test_small_highly_compressible_excel_remains_supported(self):
        comments = b'<!--' + b'X' * 1000 + b'-->'
        with zipfile.ZipFile(self.excel_path) as book:
            xml = book.read('xl/worksheets/sheet1.xml')
        self._replace_member('xl/worksheets/sheet1.xml',
                             xml.replace(b'</worksheet>', comments * 100 + b'</worksheet>'))
        result = self._merge()
        wb = load_workbook(result.path)
        try:
            self.assertEqual(wb['14시']['A1'].value, '정상 데이터')
            self.assertEqual(wb['14시']['B1'].value, '=1+2')
            self.assertEqual(wb['11시']['A2'].value, '00123')
        finally:
            wb.close()


if __name__ == '__main__':
    unittest.main()
