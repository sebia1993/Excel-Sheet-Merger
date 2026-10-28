"""Optional synthetic-only resource diagnostic, not a corporate 8h soak.

Usage (Linux): xvfb-run -a env PYTHONPATH=src python benchmarks/measure_round3_resources.py
Windows desktop: python benchmarks/measure_round3_resources.py

Measure temporary ZIP sizes just before final stitching and the exact number
of GUI log entries retained after a large synthetic burst. No user files are used.
"""
from __future__ import annotations

import os
import tempfile
import time
import tracemalloc
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from excel_sheet_merger.classify import scan_folder
from excel_sheet_merger import merge as merge_mod
from excel_sheet_merger.gui import MergeApp


def _resident_mb() -> float | None:
    # POSIX ru_maxrss is high-watermark (KiB on Linux, bytes on macOS).
    if os.name == 'nt':
        return None
    import resource
    value = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    return value / (1024 * 1024 if os.uname().sysname == 'Darwin' else 1024)


def synthetic_staging() -> None:
    with tempfile.TemporaryDirectory() as td:
        root = Path(td)
        with (root / 'A사업장_20261028_11시.csv').open('w', encoding='utf-8') as output:
            output.write('ID,문자열,메모\n')
            for i in range(20000):
                output.write(f'{i:07d},192.0.2.{i % 254 + 1},한글 업무 데이터 {i}\n')
        wb = Workbook()
        wb.active['A1'] = 'Excel'
        wb.active['B1'] = '=1+2'
        wb.save(root / 'A사업장_20261028_14시.xlsx')
        wb.close()
        original = merge_mod._stitch_csv_sheets
        stats = {}

        def inspect(base, csv, replacements, dest):
            stats['base_bytes'] = base.stat().st_size
            stats['csv_bytes'] = csv.stat().st_size
            original(base, csv, replacements, dest)
            stats['final_bytes'] = dest.stat().st_size

        started = time.perf_counter()
        with patch.object(merge_mod, '_stitch_csv_sheets', side_effect=inspect):
            merged = merge_mod.merge_group(scan_folder(root).groups[0], root / '병합결과')
        seconds = time.perf_counter() - started
        workbook = load_workbook(merged.path, read_only=True)
        try:
            assert workbook.sheetnames == ['11시', '14시']
            assert workbook['11시']['A20001'].value == '0019999'
            assert workbook['14시']['B1'].value == '=1+2'
        finally:
            workbook.close()
        print('staging:', stats, 'peak_staged_plus_output_bytes=', sum(stats.values()),
              'elapsed_s=', round(seconds, 3))
        assert not list((root / '병합결과').glob('_병합임시_*'))


def synthetic_gui_events(count: int = 20000) -> None:
    app = MergeApp()
    try:
        tracemalloc.start()
        before = _resident_mb()
        started = time.perf_counter()
        for i in range(count):
            app.events.put(('log', f'long-run-{i:05d}: synthetic test record'))
        callbacks = 0
        while not app.events.empty():
            app._consume_events()
            callbacks += 1
        elapsed = time.perf_counter() - started
        visible = app.log.get('1.0', 'end')
        assert visible.count('long-run-') == count
        assert 'long-run-00000' in visible and f'long-run-{count-1:05d}' in visible
        current, peak = tracemalloc.get_traced_memory()
        after = _resident_mb()
        tracemalloc.stop()
        print('gui_log:', {'events': count, 'callbacks': callbacks,
              'elapsed_s': round(elapsed, 3), 'python_peak_mb': round(peak/1024/1024, 2),
              'rss_highwater_before_mb': before, 'rss_highwater_after_mb': after,
              'displayed_characters': len(visible)})
    finally:
        app.destroy()


if __name__ == '__main__':
    synthetic_staging()
    synthetic_gui_events()
