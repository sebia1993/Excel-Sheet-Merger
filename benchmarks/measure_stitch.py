"""Optional, non-CI performance check: compare actual package stitching to level-6 baseline.

Run: PYTHONPATH=src python benchmarks/measure_stitch.py
Uses synthetic data; prints tradeoff. Fails if the compression optimization does not
materially improve the ZIP stitch on this host, or workbook growth exceeds 10%.
"""
from __future__ import annotations

import shutil
import statistics
import tempfile
import time
import zipfile
from copy import copy
from pathlib import Path

from openpyxl import Workbook

from excel_sheet_merger.merge import _stitch_csv_sheets


def reference_stitch(base: Path, csv: Path, destination: Path):
    with zipfile.ZipFile(base) as src, zipfile.ZipFile(csv) as csvsrc, zipfile.ZipFile(destination, 'w') as target:
        for member in src.infolist():
            copied = copy(member)
            copied._compresslevel = 6  # CPython 3.12 default DEFLATE level for this reference
            source = csvsrc if member.filename == 'xl/worksheets/sheet2.xml' else src
            filename = 'xl/worksheets/sheet1.xml' if source is csvsrc else member.filename
            with source.open(filename) as before, target.open(copied, 'w') as after:
                shutil.copyfileobj(before, after, 1024 * 1024)


def main() -> None:
    with tempfile.TemporaryDirectory() as td:
        folder = Path(td)
        base = Workbook()
        base.active['A1'] = 'Excel'
        base.create_sheet('14시')
        base.save(folder/'base.xlsx')
        base.close()
        csv = Workbook(write_only=True)
        ws = csv.create_sheet('14시')
        for i in range(30000):
            ws.append([str(i)] + ['00123-192.168.100.1', '사업장 업무자료'+str(i), 'ABC:abcdefg:abcdef123456789']*5)
        csv.save(folder/'csv.xlsx')
        csv.close()
        scores: dict[str,list[float]] = {'baseline': [], 'candidate': []}
        sizes = {}
        for round_number in range(4):
            for label in (('baseline', 'candidate') if round_number % 2 == 0 else ('candidate', 'baseline')):
                outfile = folder/f'output-{label}.xlsx'
                start = time.perf_counter()
                if label == 'baseline':
                    reference_stitch(folder/'base.xlsx', folder/'csv.xlsx', outfile)
                else:
                    _stitch_csv_sheets(folder/'base.xlsx', folder/'csv.xlsx', {2:1}, outfile)
                elapsed = time.perf_counter()-start
                with zipfile.ZipFile(outfile) as archive:
                    assert archive.testzip() is None
                scores[label].append(elapsed)
                sizes[label] = outfile.stat().st_size
        base_time = statistics.median(scores['baseline'])
        fast_time = statistics.median(scores['candidate'])
        ratio = fast_time/base_time
        size_ratio = sizes['candidate']/sizes['baseline']
        print(f'Baseline level 6: {base_time:.3f}s, {sizes["baseline"]:,} bytes')
        print(f'Current stitched: {fast_time:.3f}s, {sizes["candidate"]:,} bytes')
        print(f'Speed ratio {ratio:.3f}, size ratio {size_ratio:.3f}')
        assert ratio <= 0.85, 'Compression step has not improved enough over the original'
        assert size_ratio <= 1.10, 'Avoid significantly bigger XLSX workbooks'


if __name__ == '__main__':
    main()
